#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
AccountingAnalyzer v2 - Spreadsheet Accounting Analysis Application
Offline desktop application for accountants.

Supported input formats:
  .xlsx / .xlsm / .xltx / .xltm  (openpyxl; macro-enabled files are analyzed, never executed)
  .xls                           (xlrd)
  .xml                           (Excel 2003 SpreadsheetML)
  .csv / .tsv / .txt             (delimiter auto-detection, numeric/date coercion)
  .ods                           (native OpenDocument reader, no extra dependency)

Command-line modes:
  python compta_analyzer_en.py                       # GUI
  python compta_analyzer_en.py --self-test FOLDER    # batch-load every spreadsheet in FOLDER, report pass/fail
  python compta_analyzer_en.py --smoke-test FILE     # scripted GUI run for CI (needs a display)
  python compta_analyzer_en.py --no-restore          # start GUI without offering to restore the last session
"""

import argparse
import csv
import io
import json
import os
import sys
import zipfile
import xml.etree.ElementTree as ET
from datetime import datetime, date
from collections import defaultdict

VERSION = '2.0'

# === Optional dependencies (checked, never auto-installed) ===
try:
    import openpyxl
    from openpyxl.utils import get_column_letter
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

    def get_column_letter(idx):
        letters = ''
        while idx > 0:
            idx, rem = divmod(idx - 1, 26)
            letters = chr(65 + rem) + letters
        return letters

try:
    import xlrd
    HAS_XLRD = True
except ImportError:
    HAS_XLRD = False

try:
    from oletools.olevba import VBA_Parser
    HAS_OLETOOLS = True
except ImportError:
    HAS_OLETOOLS = False

try:
    import matplotlib
    from matplotlib.figure import Figure
    from matplotlib.backends.backend_agg import FigureCanvasAgg
    HAS_MATPLOTLIB = True
except ImportError:
    HAS_MATPLOTLIB = False

try:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import (
        SimpleDocTemplate, Table, TableStyle, Paragraph,
        Spacer, Image as RLImage, PageBreak
    )
    from reportlab.lib.enums import TA_CENTER
    HAS_REPORTLAB = True
except ImportError:
    HAS_REPORTLAB = False

try:
    import tkinter as tk
    from tkinter import ttk, filedialog, messagebox
    TK_AVAILABLE = True
except ImportError:
    TK_AVAILABLE = False


def check_dependencies(for_gui=False):
    """Return a list of missing dependencies with install hints (no auto-install)."""
    missing = []
    if not HAS_OPENPYXL:
        missing.append(('openpyxl', 'required to read .xlsx/.xlsm files'))
    if for_gui and not HAS_MATPLOTLIB:
        missing.append(('matplotlib', 'required for charts'))
    if for_gui and not HAS_REPORTLAB:
        missing.append(('reportlab', 'required for PDF export'))
    if not HAS_XLRD:
        missing.append(('xlrd', 'optional: legacy .xls support'))
    if not HAS_OLETOOLS:
        missing.append(('oletools', 'optional: VBA macro analysis'))
    return missing


# === COLORS AND STYLE ===
COLORS = {
    'bg': '#1a1a2e',
    'bg_light': '#16213e',
    'bg_card': '#0f3460',
    'accent': '#e94560',
    'accent2': '#533483',
    'text': '#ffffff',
    'text_dim': '#a0a0b0',
    'success': '#00b894',
    'warning': '#fdcb6e',
    'danger': '#e17055',
    'border': '#2d3561',
    'chart_colors': ['#e94560', '#00b894', '#fdcb6e', '#6c5ce7',
                     '#00cec9', '#ff7675', '#74b9ff', '#a29bfe',
                     '#55efc4', '#fab1a0']
}

OPENPYXL_EXTENSIONS = {'.xlsx', '.xlsm', '.xltx', '.xltm'}
DELIMITED_EXTENSIONS = {'.csv', '.tsv', '.txt'}
SUPPORTED_EXTENSIONS = (OPENPYXL_EXTENSIONS | DELIMITED_EXTENSIONS |
                        {'.xls', '.xml', '.ods'})

SPREADSHEETML_NS = 'urn:schemas-microsoft-com:office:spreadsheet'

DATE_INPUT_FORMATS = [
    '%Y-%m-%d', '%m/%d/%Y', '%d/%m/%Y', '%m-%d-%Y', '%d-%m-%Y',
    '%Y-%m-%d %H:%M:%S', '%m/%d/%Y %H:%M', '%Y/%m/%d',
]


def fmt_money(value):
    """Format a number as currency."""
    if value is None:
        return "—"
    try:
        v = float(value)
        sign = '-' if v < 0 else ''
        v = abs(v)
        return f"{sign}${v:,.2f}"
    except (ValueError, TypeError):
        return str(value)


def fmt_pct(value):
    """Format a number as percentage."""
    if value is None:
        return "—"
    try:
        return f"{float(value):.1f}%"
    except (ValueError, TypeError):
        return str(value)


def is_number(value):
    """True for int/float values that are usable in computations (bools excluded)."""
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def to_number(value):
    """Convert a cell value to float, or raise ValueError. Bools are rejected."""
    if isinstance(value, bool):
        raise ValueError('boolean')
    return float(value)


def coerce_text_cell(text):
    """Best-effort typing for text cells (CSV and fallback loaders):
    numbers (incl. $1,234.56 and (500.00) accounting negatives), dates, else str."""
    if text is None:
        return None
    s = text.strip()
    if s == '':
        return None
    cleaned = s.replace(',', '').replace('$', '').replace('€', '').replace('£', '')
    negative = False
    if cleaned.startswith('(') and cleaned.endswith(')'):
        cleaned = cleaned[1:-1]
        negative = True
    try:
        num = float(cleaned)
        if negative:
            num = -num
        return int(num) if num.is_integer() and '.' not in cleaned and 'e' not in cleaned.lower() else num
    except ValueError:
        pass
    for fmt in DATE_INPUT_FORMATS:
        try:
            return datetime.strptime(s, fmt)
        except ValueError:
            continue
    return s


# =====================================================
# CORE CLASS: ExcelAnalyzer
# =====================================================
class ExcelAnalyzer:
    """Spreadsheet accounting file analysis engine (UI-free)."""

    def __init__(self):
        self.sheets = {}
        self.file_path = None
        self.warnings = []
        self.macro_info = {'has_macros': False, 'modules': [], 'findings': [],
                           'summary': 'No file loaded.', 'error': None}

    # ---------- Loading ----------
    def load_file(self, path):
        """Load a spreadsheet file of any supported format."""
        if not os.path.isfile(path):
            raise FileNotFoundError(f"File not found: {path}")
        self.file_path = path
        self.warnings = []
        ext = os.path.splitext(path)[1].lower()

        if ext in OPENPYXL_EXTENSIONS:
            tables = self._load_openpyxl(path)
        elif ext == '.xls':
            tables = self._load_xls(path)
        elif ext == '.xml':
            tables = self._load_spreadsheetml(path)
        elif ext in DELIMITED_EXTENSIONS:
            tables = self._load_delimited(path)
        elif ext == '.ods':
            tables = self._load_ods(path)
        else:
            tables = self._load_unknown(path)

        self.sheets = {}
        for name, data in tables.items():
            self.sheets[name] = {
                'data': data,
                'headers': data[0] if data else [],
                'rows': data[1:] if len(data) > 1 else [],
                'num_rows': len(data) - 1 if data else 0,
                'num_cols': max((len(r) for r in data), default=0)
            }
        self.macro_info = self.analyze_macros(path)
        return self.sheets

    def _load_openpyxl(self, path):
        if not HAS_OPENPYXL:
            raise RuntimeError("openpyxl is required for .xlsx/.xlsm files "
                               "(pip install openpyxl)")
        wb_values = openpyxl.load_workbook(path, data_only=True, read_only=True)
        wb_formulas = openpyxl.load_workbook(path, data_only=False, read_only=True)
        tables = {}
        for name in wb_values.sheetnames:
            ws_v = wb_values[name]
            ws_f = wb_formulas[name]
            data = [list(row) for row in ws_v.iter_rows(values_only=True)]
            uncached = 0
            for row_v, row_f in zip(data, ws_f.iter_rows(values_only=True)):
                for v, f in zip(row_v, row_f):
                    if v is None and isinstance(f, str) and f.startswith('='):
                        uncached += 1
            if uncached:
                self.warnings.append(
                    f"Sheet '{name}': {uncached} formula cell(s) have no cached value "
                    f"(file never recalculated by a spreadsheet app) — they load as empty.")
            tables[name] = data
        wb_values.close()
        wb_formulas.close()
        return tables

    def _load_xls(self, path):
        if not HAS_XLRD:
            raise RuntimeError("xlrd is required for legacy .xls files "
                               "(pip install xlrd)")
        book = xlrd.open_workbook(path)
        tables = {}
        for sheet in book.sheets():
            data = []
            for r in range(sheet.nrows):
                row = []
                for c in range(sheet.ncols):
                    cell = sheet.cell(r, c)
                    if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                        row.append(None)
                    elif cell.ctype == xlrd.XL_CELL_DATE:
                        row.append(xlrd.xldate.xldate_as_datetime(cell.value, book.datemode))
                    elif cell.ctype == xlrd.XL_CELL_NUMBER:
                        v = cell.value
                        row.append(int(v) if float(v).is_integer() else v)
                    elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                        row.append(bool(cell.value))
                    elif cell.ctype == xlrd.XL_CELL_ERROR:
                        row.append(None)
                    else:
                        row.append(cell.value)
                data.append(row)
            tables[sheet.name] = data
        return tables

    def _load_spreadsheetml(self, path):
        """Excel 2003 SpreadsheetML (.xml) reader."""
        tree = ET.parse(path)
        root = tree.getroot()
        if not root.tag.endswith('}Workbook') and root.tag != 'Workbook':
            raise RuntimeError("Not a SpreadsheetML (Excel 2003 XML) workbook.")
        ns = {'ss': SPREADSHEETML_NS}
        tables = {}
        for wsi, worksheet in enumerate(root.findall('ss:Worksheet', ns), start=1):
            name = worksheet.get(f'{{{SPREADSHEETML_NS}}}Name') or f'Sheet{wsi}'
            data = []
            table = worksheet.find('ss:Table', ns)
            if table is None:
                tables[name] = data
                continue
            row_cursor = 0
            for row_el in table.findall('ss:Row', ns):
                row_index = row_el.get(f'{{{SPREADSHEETML_NS}}}Index')
                if row_index:
                    while row_cursor < int(row_index) - 1:
                        data.append([])
                        row_cursor += 1
                row = []
                col_cursor = 0
                for cell_el in row_el.findall('ss:Cell', ns):
                    cell_index = cell_el.get(f'{{{SPREADSHEETML_NS}}}Index')
                    if cell_index:
                        while col_cursor < int(cell_index) - 1:
                            row.append(None)
                            col_cursor += 1
                    data_el = cell_el.find('ss:Data', ns)
                    row.append(self._parse_ssml_value(data_el))
                    col_cursor += 1
                data.append(row)
                row_cursor += 1
            tables[name] = data
        if not tables:
            raise RuntimeError("SpreadsheetML workbook contains no worksheets.")
        return tables

    @staticmethod
    def _parse_ssml_value(data_el):
        if data_el is None or data_el.text is None:
            return None
        text = data_el.text
        ss_type = data_el.get(f'{{{SPREADSHEETML_NS}}}Type', 'String')
        if ss_type == 'Number':
            v = float(text)
            return int(v) if v.is_integer() else v
        if ss_type == 'DateTime':
            try:
                return datetime.fromisoformat(text.rstrip('Z'))
            except ValueError:
                return text
        if ss_type == 'Boolean':
            return bool(int(text))
        return text

    def _load_delimited(self, path):
        """CSV/TSV/TXT with delimiter sniffing and numeric/date coercion."""
        with open(path, 'r', encoding='utf-8-sig', errors='replace', newline='') as f:
            sample = f.read(64 * 1024)
            f.seek(0)
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=',;\t|')
            except csv.Error:
                class dialect(csv.excel):
                    delimiter = '\t' if path.lower().endswith('.tsv') else ','
            reader = csv.reader(f, dialect)
            data = [[coerce_text_cell(cell) for cell in row] for row in reader]
        name = os.path.splitext(os.path.basename(path))[0] or 'Sheet1'
        return {name: data}

    def _load_ods(self, path):
        """Native OpenDocument Spreadsheet reader (content.xml, no odfpy needed)."""
        NS_TABLE = 'urn:oasis:names:tc:opendocument:xmlns:table:1.0'
        NS_OFFICE = 'urn:oasis:names:tc:opendocument:xmlns:office:1.0'
        with zipfile.ZipFile(path) as zf:
            with zf.open('content.xml') as f:
                root = ET.parse(f).getroot()
        tables = {}
        for ti, table in enumerate(root.iter(f'{{{NS_TABLE}}}table'), start=1):
            name = table.get(f'{{{NS_TABLE}}}name') or f'Sheet{ti}'
            data = []
            for row_el in table.iter(f'{{{NS_TABLE}}}table-row'):
                row_repeat = int(row_el.get(f'{{{NS_TABLE}}}number-rows-repeated', '1'))
                row = []
                for cell_el in row_el:
                    if not cell_el.tag.endswith('}table-cell'):
                        continue
                    repeat = int(cell_el.get(f'{{{NS_TABLE}}}number-columns-repeated', '1'))
                    vtype = cell_el.get(f'{{{NS_OFFICE}}}value-type')
                    if vtype in ('float', 'currency', 'percentage'):
                        v = float(cell_el.get(f'{{{NS_OFFICE}}}value', '0'))
                        value = int(v) if v.is_integer() else v
                    elif vtype == 'date':
                        raw = cell_el.get(f'{{{NS_OFFICE}}}date-value', '')
                        try:
                            value = datetime.fromisoformat(raw)
                        except ValueError:
                            value = raw or None
                    elif vtype == 'boolean':
                        value = cell_el.get(f'{{{NS_OFFICE}}}boolean-value') == 'true'
                    elif vtype is not None:
                        texts = [t for t in cell_el.itertext()]
                        value = ''.join(texts) if texts else None
                    else:
                        value = None
                    if value is None and repeat > 64:
                        # trailing padding emitted by some writers — ignore
                        continue
                    row.extend([value] * repeat)
                while row and row[-1] is None:
                    row.pop()
                if row_repeat > 64 and not any(v is not None for v in row):
                    continue
                for _ in range(row_repeat):
                    data.append(list(row))
            while data and not any(v is not None for v in data[-1]):
                data.pop()
            tables[name] = data
        if not tables:
            raise RuntimeError("ODS file contains no sheets.")
        return tables

    def _load_unknown(self, path):
        """Unknown extension: try openpyxl, then delimited text."""
        errors = []
        if HAS_OPENPYXL:
            try:
                return self._load_openpyxl(path)
            except Exception as e:
                errors.append(f"openpyxl: {e}")
        try:
            return self._load_delimited(path)
        except Exception as e:
            errors.append(f"text: {e}")
        supported = ', '.join(sorted(SUPPORTED_EXTENSIONS))
        raise RuntimeError(
            f"Unsupported file format. Supported: {supported}. Tried: {'; '.join(errors)}")

    # ---------- Macro analysis (static only — macros are NEVER executed) ----------
    def analyze_macros(self, path):
        """Static VBA macro inspection: presence, module names, olevba findings."""
        info = {'has_macros': False, 'modules': [], 'findings': [],
                'summary': '', 'error': None}
        ext = os.path.splitext(path)[1].lower()

        if ext in OPENPYXL_EXTENSIONS or ext == '.ods':
            try:
                with zipfile.ZipFile(path) as zf:
                    names = zf.namelist()
                info['has_macros'] = any(n.lower().endswith('vbaproject.bin') for n in names)
                if ext == '.ods' and any(n.startswith('Basic/') for n in names):
                    info['has_macros'] = True
            except zipfile.BadZipFile:
                pass
        elif ext == '.xls':
            info['has_macros'] = None  # unknown until olevba runs

        if HAS_OLETOOLS and ext not in DELIMITED_EXTENSIONS and ext != '.xml':
            try:
                parser = VBA_Parser(path)
                try:
                    detected = parser.detect_vba_macros()
                    # zip-level vbaProject.bin presence always wins: a macro-enabled
                    # workbook stays flagged even if the VBA project can't be parsed
                    info['has_macros'] = detected or bool(info['has_macros'])
                    if detected:
                        seen = set()
                        for _fn, _stream, vba_filename, _code in parser.extract_macros():
                            if vba_filename not in seen:
                                seen.add(vba_filename)
                                info['modules'].append(vba_filename)
                        for kw_type, keyword, description in parser.analyze_macros():
                            info['findings'].append({
                                'type': str(kw_type), 'keyword': str(keyword),
                                'description': str(description)})
                finally:
                    parser.close()
            except Exception as e:
                if info['has_macros']:
                    info['error'] = f"VBA project present but could not be parsed: {e}"
                # non-OLE/zip inputs (plain csv-like) simply have no macros
        elif info['has_macros'] and not HAS_OLETOOLS:
            info['error'] = ("oletools not installed — macro presence detected but "
                             "content not analyzed (pip install oletools)")

        if ext == '.xls' and info['has_macros'] is None:
            info['has_macros'] = False

        n_susp = sum(1 for f in info['findings'] if 'susp' in f['type'].lower())
        n_auto = sum(1 for f in info['findings'] if 'auto' in f['type'].lower())
        if not info['has_macros']:
            info['summary'] = 'No VBA macros detected.'
        else:
            parts = [f"VBA macros detected ({len(info['modules'])} module(s))"]
            if n_auto:
                parts.append(f"{n_auto} auto-exec trigger(s)")
            if n_susp:
                parts.append(f"{n_susp} suspicious keyword(s)")
            info['summary'] = ', '.join(parts) + '. Macros were analyzed statically, never executed.'
        return info

    # ---------- Introspection ----------
    def get_sheet_names(self):
        return list(self.sheets.keys())

    def get_sheet_data(self, sheet_name):
        return self.sheets.get(sheet_name, {})

    def detect_numeric_columns(self, sheet_name):
        """Detect numeric columns (index, header) — booleans don't count as numeric."""
        sheet = self.sheets.get(sheet_name, {})
        rows = sheet.get('rows', [])
        headers = sheet.get('headers', [])
        numeric_cols = []
        for i, h in enumerate(headers):
            count = 0
            total = 0
            for row in rows:
                if i < len(row) and row[i] is not None:
                    total += 1
                    try:
                        to_number(row[i])
                        count += 1
                    except (ValueError, TypeError):
                        pass
            if total > 0 and count / total > 0.5:
                numeric_cols.append((i, h))
        return numeric_cols

    def detect_date_columns(self, sheet_name):
        """Detect date columns (index, header)."""
        sheet = self.sheets.get(sheet_name, {})
        rows = sheet.get('rows', [])
        headers = sheet.get('headers', [])
        date_cols = []
        for i, h in enumerate(headers):
            count = 0
            total = 0
            for row in rows:
                if i < len(row) and row[i] is not None:
                    total += 1
                    if isinstance(row[i], (datetime, date)):
                        count += 1
            if total > 0 and count / total > 0.5:
                date_cols.append((i, h))
        return date_cols

    def detect_category_columns(self, sheet_name, max_unique=50):
        """Detect text columns suitable for a category breakdown."""
        sheet = self.sheets.get(sheet_name, {})
        rows = sheet.get('rows', [])
        headers = sheet.get('headers', [])
        numeric = {i for i, _ in self.detect_numeric_columns(sheet_name)}
        dates = {i for i, _ in self.detect_date_columns(sheet_name)}
        result = []
        for i, h in enumerate(headers):
            if i in numeric or i in dates:
                continue
            values = {str(row[i]).strip() for row in rows
                      if i < len(row) and row[i] is not None}
            if 2 <= len(values) <= max_unique:
                result.append((i, h))
        return result

    # ---------- Analyses ----------
    def analyze_financial_summary(self, sheet_name):
        """Financial summary analysis of a sheet."""
        sheet = self.sheets.get(sheet_name, {})
        rows = sheet.get('rows', [])
        numeric_cols = self.detect_numeric_columns(sheet_name)

        results = {}
        for col_idx, col_name in numeric_cols:
            values = []
            for row in rows:
                if col_idx < len(row) and row[col_idx] is not None:
                    try:
                        values.append(to_number(row[col_idx]))
                    except (ValueError, TypeError):
                        pass
            if values:
                results[col_name] = {
                    'total': sum(values),
                    'average': sum(values) / len(values),
                    'min': min(values),
                    'max': max(values),
                    'count': len(values),
                    'positives': sum(1 for v in values if v > 0),
                    'negatives': sum(1 for v in values if v < 0),
                    'total_positives': sum(v for v in values if v > 0),
                    'total_negatives': sum(v for v in values if v < 0),
                }
        return results

    def analyze_by_category(self, sheet_name, cat_col_idx, val_col_idx):
        """Analysis by category (breakdown)."""
        sheet = self.sheets.get(sheet_name, {})
        rows = sheet.get('rows', [])
        categories = defaultdict(lambda: {'total': 0, 'count': 0, 'values': []})

        for row in rows:
            if cat_col_idx < len(row) and val_col_idx < len(row):
                cat = row[cat_col_idx]
                val = row[val_col_idx]
                if cat is not None and val is not None:
                    try:
                        v = to_number(val)
                    except (ValueError, TypeError):
                        continue
                    cat_str = str(cat).strip()
                    categories[cat_str]['total'] += v
                    categories[cat_str]['count'] += 1
                    categories[cat_str]['values'].append(v)

        for cat in categories:
            vals = categories[cat]['values']
            categories[cat]['average'] = categories[cat]['total'] / categories[cat]['count']
            categories[cat]['min'] = min(vals)
            categories[cat]['max'] = max(vals)

        return dict(categories)

    def analyze_time_series(self, sheet_name, date_col_idx, val_col_idx):
        """Time series analysis: [(datetime, value)] sorted chronologically."""
        sheet = self.sheets.get(sheet_name, {})
        rows = sheet.get('rows', [])
        series = []

        for row in rows:
            if date_col_idx < len(row) and val_col_idx < len(row):
                date_val = row[date_col_idx]
                num_val = row[val_col_idx]
                if date_val is None or num_val is None:
                    continue
                try:
                    v = to_number(num_val)
                except (ValueError, TypeError):
                    continue
                if isinstance(date_val, datetime):
                    series.append((date_val, v))
                elif isinstance(date_val, date):
                    series.append((datetime(date_val.year, date_val.month, date_val.day), v))
                elif isinstance(date_val, str):
                    for fmt in DATE_INPUT_FORMATS:
                        try:
                            series.append((datetime.strptime(date_val, fmt), v))
                            break
                        except ValueError:
                            continue

        series.sort(key=lambda x: x[0])
        return series

    def detect_anomalies(self, sheet_name, col_idx, threshold=2.0):
        """Detect anomalies in a column (deviation > threshold * std dev)."""
        sheet = self.sheets.get(sheet_name, {})
        rows = sheet.get('rows', [])
        values = []

        for i, row in enumerate(rows):
            if col_idx < len(row) and row[col_idx] is not None:
                try:
                    values.append((i, to_number(row[col_idx])))
                except (ValueError, TypeError):
                    pass

        if len(values) < 3:
            return []

        nums = [v for _, v in values]
        mean = sum(nums) / len(nums)
        variance = sum((x - mean) ** 2 for x in nums) / len(nums)
        std = variance ** 0.5

        if std == 0:
            return []

        anomalies = []
        for idx, val in values:
            z_score = abs(val - mean) / std
            if z_score > threshold:
                anomalies.append({
                    'row': idx + 2,
                    'value': val,
                    'z_score': z_score,
                    'mean': mean,
                    'deviation': val - mean
                })
        return anomalies

    def bank_reconciliation(self, sheet1_name, col1_idx, sheet2_name, col2_idx, tolerance=0.01):
        """Reconcile two amount columns. Each entry matches at most once, so
        duplicate amounts reconcile one-to-one (sorted two-pointer, O(n log n))."""
        def extract(sheet_name, col_idx):
            vals = []
            for i, row in enumerate(self.sheets.get(sheet_name, {}).get('rows', [])):
                if col_idx < len(row) and row[col_idx] is not None:
                    try:
                        vals.append((i, to_number(row[col_idx])))
                    except (ValueError, TypeError):
                        pass
            return vals

        vals1 = extract(sheet1_name, col1_idx)
        vals2 = extract(sheet2_name, col2_idx)

        order1 = sorted(range(len(vals1)), key=lambda k: vals1[k][1])
        order2 = sorted(range(len(vals2)), key=lambda k: vals2[k][1])

        matched = []
        used1, used2 = set(), set()
        i = j = 0
        while i < len(order1) and j < len(order2):
            idx1, v1 = vals1[order1[i]]
            idx2, v2 = vals2[order2[j]]
            if abs(v1 - v2) <= tolerance:
                matched.append({'sheet1_row': idx1 + 2, 'sheet2_row': idx2 + 2, 'value': v1})
                used1.add(order1[i])
                used2.add(order2[j])
                i += 1
                j += 1
            elif v1 < v2:
                i += 1
            else:
                j += 1

        unmatched1 = [vals1[k] for k in range(len(vals1)) if k not in used1]
        unmatched2 = [vals2[k] for k in range(len(vals2)) if k not in used2]

        return {
            'matched': matched,
            'unmatched_sheet1': [{'row': idx + 2, 'value': v} for idx, v in unmatched1],
            'unmatched_sheet2': [{'row': idx + 2, 'value': v} for idx, v in unmatched2],
            'total_sheet1': sum(v for _, v in vals1),
            'total_sheet2': sum(v for _, v in vals2),
            'difference': sum(v for _, v in vals1) - sum(v for _, v in vals2),
            'match_rate': len(matched) / max(len(vals1), len(vals2), 1) * 100
        }

    def compute_ratios(self, summary):
        """Compute basic financial ratios from summary."""
        ratios = {}
        keys = list(summary.keys())

        rev_keys = [k for k in keys if any(w in str(k).lower() for w in
                    ['revenue', 'sales', 'income', 'receipts', 'turnover'])]
        exp_keys = [k for k in keys if any(w in str(k).lower() for w in
                    ['expense', 'cost', 'charge', 'purchase', 'spending'])]

        if rev_keys and exp_keys:
            revenue = sum(summary[k]['total'] for k in rev_keys)
            expenses = abs(sum(summary[k]['total'] for k in exp_keys))
            if revenue != 0:
                ratios['Margin (%)'] = ((revenue - expenses) / revenue) * 100
                ratios['Expense/Revenue Ratio'] = (expenses / revenue) * 100

        for k in keys:
            s = summary[k]
            if s['count'] > 0:
                ratios[f"Average {k}"] = s['average']
                if s['total_positives'] != 0:
                    ratios[f"Neg/Pos Ratio {k}"] = abs(s['total_negatives']) / s['total_positives'] * 100

        return ratios


# =====================================================
# SESSION AUTOSAVE
# =====================================================
class SessionManager:
    """Persists application state to ~/.accounting_analyzer/session.json."""

    def __init__(self, directory=None):
        self.directory = directory or os.path.join(
            os.path.expanduser('~'), '.accounting_analyzer')
        self.session_path = os.path.join(self.directory, 'session.json')

    def save(self, state):
        """Atomic save (write temp file, then rename)."""
        state = dict(state)
        state['saved_at'] = datetime.now().isoformat(timespec='seconds')
        state['version'] = VERSION
        os.makedirs(self.directory, exist_ok=True)
        tmp_path = self.session_path + '.tmp'
        with open(tmp_path, 'w', encoding='utf-8') as f:
            json.dump(state, f, indent=2)
        os.replace(tmp_path, self.session_path)
        return state['saved_at']

    def load(self):
        try:
            with open(self.session_path, 'r', encoding='utf-8') as f:
                return json.load(f)
        except (OSError, ValueError):
            return None

    def clear(self):
        try:
            os.remove(self.session_path)
        except OSError:
            pass


# =====================================================
# PDF REPORT
# =====================================================
if HAS_REPORTLAB:
    class PDFReportGenerator:
        """Generates a complete PDF report."""

        def __init__(self, file_path):
            self.file_path = file_path
            self.styles = getSampleStyleSheet()
            self._setup_styles()
            self.elements = []

        def _setup_styles(self):
            self.styles.add(ParagraphStyle(
                'CustomTitle',
                parent=self.styles['Title'],
                fontSize=24,
                spaceAfter=30,
                textColor=colors.HexColor('#1a1a2e'),
                alignment=TA_CENTER
            ))
            self.styles.add(ParagraphStyle(
                'SectionTitle',
                parent=self.styles['Heading2'],
                fontSize=14,
                spaceBefore=20,
                spaceAfter=10,
                textColor=colors.HexColor('#0f3460'),
                borderWidth=1,
                borderColor=colors.HexColor('#e94560'),
                borderPadding=5
            ))
            self.styles.add(ParagraphStyle(
                'CustomBody',
                parent=self.styles['Normal'],
                fontSize=10,
                spaceAfter=6,
                leading=14
            ))

        def add_title(self, title, subtitle=""):
            self.elements.append(Spacer(1, 2 * cm))
            self.elements.append(Paragraph(title, self.styles['CustomTitle']))
            if subtitle:
                self.elements.append(Paragraph(subtitle, self.styles['CustomBody']))
            self.elements.append(Spacer(1, 1 * cm))
            line_table = Table([['']], colWidths=[17 * cm])
            line_table.setStyle(TableStyle([
                ('LINEBELOW', (0, 0), (-1, -1), 2, colors.HexColor('#e94560')),
            ]))
            self.elements.append(line_table)
            self.elements.append(Spacer(1, 0.5 * cm))

        def add_section(self, title):
            self.elements.append(Spacer(1, 0.5 * cm))
            self.elements.append(Paragraph(title, self.styles['SectionTitle']))

        def add_text(self, text):
            self.elements.append(Paragraph(text, self.styles['CustomBody']))

        def add_table(self, headers, rows, col_widths=None):
            data = [headers] + rows
            if col_widths is None:
                col_widths = [17 * cm / len(headers)] * len(headers)

            table = Table(data, colWidths=col_widths)
            table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0f3460')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
                ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, 0), 10),
                ('FONTSIZE', (0, 1), (-1, -1), 9),
                ('BOTTOMPADDING', (0, 0), (-1, 0), 10),
                ('TOPPADDING', (0, 0), (-1, 0), 10),
                ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#cccccc')),
                ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f5f5f5')]),
                ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ]))
            self.elements.append(table)
            self.elements.append(Spacer(1, 0.3 * cm))

        def add_chart_image(self, fig):
            """Add a matplotlib chart to the PDF."""
            FigureCanvasAgg(fig)
            buf = io.BytesIO()
            fig.savefig(buf, format='png', dpi=150, bbox_inches='tight',
                        facecolor='white', edgecolor='none')
            buf.seek(0)
            img = RLImage(buf, width=16 * cm, height=9 * cm)
            self.elements.append(img)
            self.elements.append(Spacer(1, 0.5 * cm))

        def add_kpi_row(self, kpis):
            """Add a row of KPIs (dict {label: value})."""
            headers = list(kpis.keys())
            values = [str(v) for v in kpis.values()]
            data = [headers, values]
            col_w = [17 * cm / len(headers)] * len(headers)
            table = Table(data, colWidths=col_w)
            table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#533483')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
                ('TEXTCOLOR', (0, 1), (-1, 1), colors.HexColor('#1a1a2e')),
                ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, 0), 9),
                ('FONTSIZE', (0, 1), (-1, 1), 12),
                ('FONTNAME', (0, 1), (-1, 1), 'Helvetica-Bold'),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 10),
                ('TOPPADDING', (0, 0), (-1, -1), 10),
                ('BOX', (0, 0), (-1, -1), 1, colors.HexColor('#533483')),
                ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#dddddd')),
            ]))
            self.elements.append(table)
            self.elements.append(Spacer(1, 0.3 * cm))

        def add_page_break(self):
            self.elements.append(PageBreak())

        def generate(self):
            doc = SimpleDocTemplate(
                self.file_path,
                pagesize=A4,
                rightMargin=1.5 * cm,
                leftMargin=1.5 * cm,
                topMargin=1.5 * cm,
                bottomMargin=1.5 * cm
            )

            def add_footer(canvas, doc):
                canvas.saveState()
                canvas.setFont('Helvetica', 8)
                canvas.setFillColor(colors.gray)
                canvas.drawString(1.5 * cm, 1 * cm,
                                  f"AccountingAnalyzer — Report generated on "
                                  f"{datetime.now().strftime('%m/%d/%Y at %H:%M')}")
                canvas.drawRightString(A4[0] - 1.5 * cm, 1 * cm, f"Page {doc.page}")
                canvas.restoreState()

            doc.build(self.elements, onFirstPage=add_footer, onLaterPages=add_footer)


def build_pdf_report(analyzer, path, include=None, anomaly_threshold=2.0):
    """Build the full PDF report from an analyzer. UI-free, reused by the GUI,
    the smoke test and callers embedding the engine.

    include: set of section names among {'summary', 'categories', 'anomalies',
    'charts', 'macros'} (default: all)."""
    if not HAS_REPORTLAB:
        raise RuntimeError("reportlab is required for PDF export (pip install reportlab)")
    if include is None:
        include = {'summary', 'categories', 'anomalies', 'charts', 'macros'}

    pdf = PDFReportGenerator(path)
    source_name = os.path.basename(analyzer.file_path) if analyzer.file_path else "File"
    pdf.add_title(
        "Accounting Analysis Report",
        f"Source: {source_name} — Generated on {datetime.now().strftime('%m/%d/%Y at %H:%M')}"
    )

    if analyzer.warnings:
        pdf.add_section("Data Loading Warnings")
        for w in analyzer.warnings:
            pdf.add_text(f"• {w}")

    if 'macros' in include:
        info = analyzer.macro_info
        pdf.add_section("Macro Analysis")
        pdf.add_text(info['summary'] or 'No macro information.')
        if info.get('error'):
            pdf.add_text(f"<b>Note:</b> {info['error']}")
        if info['modules']:
            pdf.add_text("<b>Modules:</b> " + ", ".join(info['modules'][:15]))
        if info['findings']:
            rows = [[f['type'][:18], f['keyword'][:24], f['description'][:60]]
                    for f in info['findings'][:10]]
            pdf.add_table(['Type', 'Keyword', 'Description'], rows,
                          col_widths=[3.5 * cm, 4.5 * cm, 9 * cm])

    if 'summary' in include:
        for sheet_name in analyzer.get_sheet_names():
            summary = analyzer.analyze_financial_summary(sheet_name)
            if not summary:
                continue

            pdf.add_section(f"Financial Summary — {sheet_name}")

            kpis = {}
            for col, stats in list(summary.items())[:4]:
                kpis[str(col)[:15]] = fmt_money(stats['total'])
            if kpis:
                pdf.add_kpi_row(kpis)

            headers = ['Column', 'Total', 'Average', 'Min', 'Max', 'Count']
            rows = []
            for col, stats in summary.items():
                rows.append([
                    str(col)[:20],
                    fmt_money(stats['total']),
                    fmt_money(stats['average']),
                    fmt_money(stats['min']),
                    fmt_money(stats['max']),
                    str(stats['count'])
                ])
            pdf.add_table(headers, rows)

            ratios = analyzer.compute_ratios(summary)
            if ratios:
                ratio_kpis = {k[:18]: fmt_pct(v) if '%' in k or 'ratio' in k.lower() else fmt_money(v)
                              for k, v in list(ratios.items())[:4]}
                pdf.add_text("<b>Computed Ratios:</b>")
                pdf.add_kpi_row(ratio_kpis)

    if 'categories' in include:
        wrote_header = False
        for sheet_name in analyzer.get_sheet_names():
            cat_cols = analyzer.detect_category_columns(sheet_name)
            num_cols = analyzer.detect_numeric_columns(sheet_name)
            if not cat_cols or not num_cols:
                continue
            cat_idx, cat_name = cat_cols[0]
            val_idx, val_name = num_cols[0]
            results = analyzer.analyze_by_category(sheet_name, cat_idx, val_idx)
            if not results:
                continue
            if not wrote_header:
                pdf.add_page_break()
                wrote_header = True
            pdf.add_section(f"Breakdown — {sheet_name} ({cat_name} × {val_name})")
            sorted_results = sorted(results.items(), key=lambda x: abs(x[1]['total']),
                                    reverse=True)
            grand_total = sum(abs(v['total']) for v in results.values())
            rows = []
            for cat, stats in sorted_results[:15]:
                pct = (abs(stats['total']) / grand_total * 100) if grand_total else 0
                rows.append([str(cat)[:24], fmt_money(stats['total']),
                             fmt_money(stats['average']), str(stats['count']), fmt_pct(pct)])
            pdf.add_table(['Category', 'Total', 'Average', 'Count', 'Share %'], rows)

    if 'anomalies' in include:
        pdf.add_page_break()
        pdf.add_section(f"Anomaly Detection (threshold = {anomaly_threshold}σ)")
        found_any = False
        for sheet_name in analyzer.get_sheet_names():
            for col_idx, col_name in analyzer.detect_numeric_columns(sheet_name):
                anomalies = analyzer.detect_anomalies(sheet_name, col_idx, anomaly_threshold)
                if anomalies:
                    found_any = True
                    pdf.add_text(f"<b>{sheet_name} — {col_name}</b>: "
                                 f"{len(anomalies)} anomaly(ies)")
                    rows = [[str(a['row']), fmt_money(a['value']),
                             fmt_money(a['mean']), fmt_money(a['deviation']),
                             f"{a['z_score']:.2f}"] for a in anomalies[:10]]
                    pdf.add_table(['Row', 'Value', 'Mean', 'Deviation', 'Z-Score'], rows)
        if not found_any:
            pdf.add_text(f"No significant anomalies detected (threshold = {anomaly_threshold}σ).")

    if 'charts' in include and HAS_MATPLOTLIB:
        wrote_header = False
        for sheet_name in analyzer.get_sheet_names():
            summary = analyzer.analyze_financial_summary(sheet_name)
            if not summary:
                continue
            if not wrote_header:
                pdf.add_page_break()
                pdf.add_section("Charts")
                wrote_header = True
            fig = Figure(figsize=(8, 3.5), dpi=100)
            ax = fig.add_subplot(111)
            names = [str(k)[:15] for k in summary.keys()]
            totals = [summary[k]['total'] for k in summary.keys()]
            colors_bar = ['#00b894' if v >= 0 else '#e17055' for v in totals]
            ax.bar(names, totals, color=colors_bar)
            ax.set_title(f"Totals — {sheet_name}", fontsize=11)
            ax.tick_params(axis='x', rotation=30, labelsize=8)
            fig.tight_layout()
            pdf.add_chart_image(fig)

    pdf.generate()
    return path


# =====================================================
# SELF-TEST (headless batch validation of a folder or file)
# =====================================================
def run_self_test(target):
    """Batch-load every supported spreadsheet under `target`, print a report,
    return 0 if every file loaded successfully."""
    if os.path.isfile(target):
        files = [target]
    elif os.path.isdir(target):
        files = []
        for root, _dirs, names in os.walk(target):
            for n in sorted(names):
                if os.path.splitext(n)[1].lower() in SUPPORTED_EXTENSIONS:
                    files.append(os.path.join(root, n))
    else:
        print(f"Path not found: {target}")
        return 2

    if not files:
        print(f"No supported spreadsheet files found under: {target}")
        print(f"Supported extensions: {', '.join(sorted(SUPPORTED_EXTENSIONS))}")
        return 2

    print(f"AccountingAnalyzer v{VERSION} self-test — {len(files)} file(s)\n")
    failures = 0
    for path in files:
        rel = os.path.relpath(path, target if os.path.isdir(target) else os.path.dirname(target) or '.')
        analyzer = ExcelAnalyzer()
        try:
            sheets = analyzer.load_file(path)
            total_rows = sum(s['num_rows'] for s in sheets.values())
            n_numeric = sum(len(analyzer.detect_numeric_columns(s)) for s in sheets)
            macro = 'macros: YES' if analyzer.macro_info['has_macros'] else 'macros: no'
            print(f"  PASS  {rel}")
            print(f"        {len(sheets)} sheet(s), {total_rows} data row(s), "
                  f"{n_numeric} numeric column(s), {macro}")
            if analyzer.macro_info['has_macros']:
                print(f"        {analyzer.macro_info['summary']}")
            for w in analyzer.warnings:
                print(f"        warning: {w}")
        except Exception as e:
            failures += 1
            print(f"  FAIL  {rel}")
            print(f"        {type(e).__name__}: {e}")

    print(f"\n{len(files) - failures}/{len(files)} file(s) loaded successfully.")
    if failures:
        print("Some files failed — see messages above.")
    return 1 if failures else 0


# =====================================================
# TKINTER APPLICATION
# =====================================================
if TK_AVAILABLE:
    class AccountingApp(tk.Tk):
        """Main application."""

        FILE_TYPES = [
            ('All spreadsheets', '*.xlsx *.xlsm *.xltx *.xltm *.xls *.xml *.csv *.tsv *.txt *.ods'),
            ('Excel', '*.xlsx *.xlsm *.xltx *.xltm *.xls'),
            ('Excel 2003 XML', '*.xml'),
            ('CSV / TSV', '*.csv *.tsv *.txt'),
            ('OpenDocument', '*.ods'),
            ('All files', '*.*')
        ]

        def __init__(self, restore_session=True):
            super().__init__()
            self.title(f"AccountingAnalyzer v{VERSION} — Spreadsheet Analysis Tool")
            self.geometry("1280x800")
            self.minsize(1024, 700)
            self.configure(bg=COLORS['bg'])

            self.analyzer = ExcelAnalyzer()
            self.session = SessionManager()
            self.current_sheet = None
            self.page_figures = {}

            self._setup_styles()
            self._build_ui()

            self.protocol('WM_DELETE_WINDOW', self._on_close)
            self.after(60000, self._autosave_tick)

            if restore_session:
                self.after(200, self._offer_session_restore)

        # ---------- styles ----------
        def _setup_styles(self):
            style = ttk.Style()
            style.theme_use('clam')

            style.configure('TFrame', background=COLORS['bg'])
            style.configure('Card.TFrame', background=COLORS['bg_card'])
            style.configure('TLabel', background=COLORS['bg'], foreground=COLORS['text'],
                            font=('Segoe UI', 10))
            style.configure('Title.TLabel', font=('Segoe UI', 20, 'bold'),
                            foreground=COLORS['text'])
            style.configure('Subtitle.TLabel', font=('Segoe UI', 12),
                            foreground=COLORS['text_dim'])
            style.configure('Section.TLabel', font=('Segoe UI', 14, 'bold'),
                            foreground=COLORS['text'])

            style.configure('Accent.TButton', background=COLORS['accent'],
                            foreground='white', font=('Segoe UI', 11, 'bold'),
                            padding=(20, 10))
            style.map('Accent.TButton',
                      background=[('active', '#c0392b'), ('disabled', '#555')])

            style.configure('Secondary.TButton', background=COLORS['accent2'],
                            foreground='white', font=('Segoe UI', 10),
                            padding=(15, 8))
            style.map('Secondary.TButton',
                      background=[('active', '#6c3483')])

            style.configure('Nav.TButton', background=COLORS['bg_light'],
                            foreground=COLORS['text'], font=('Segoe UI', 10),
                            padding=(15, 10))
            style.map('Nav.TButton',
                      background=[('active', COLORS['bg_card'])])

            style.configure('Treeview', background=COLORS['bg_light'],
                            foreground=COLORS['text'], fieldbackground=COLORS['bg_light'],
                            font=('Segoe UI', 9), rowheight=28)
            style.configure('Treeview.Heading', background=COLORS['bg_card'],
                            foreground=COLORS['text'], font=('Segoe UI', 10, 'bold'))
            style.map('Treeview', background=[('selected', COLORS['accent2'])])

            style.configure('TCombobox', fieldbackground=COLORS['bg_light'],
                            background=COLORS['bg_card'], foreground=COLORS['text'])

        # ---------- shared widget helpers ----------
        @staticmethod
        def _column_labels(headers):
            return [f"{get_column_letter(i + 1)} — "
                    f"{str(h) if h not in (None, '') else 'Col ' + str(i + 1)}"
                    for i, h in enumerate(headers)]

        def _set_column_choices(self, combo, headers, indices=None):
            """Fill a combobox with disambiguated column labels. `indices` restricts
            the choices to a subset of column indices (e.g. numeric columns only)."""
            if indices is None:
                indices = list(range(len(headers)))
            labels = self._column_labels(headers)
            combo['values'] = [labels[i] for i in indices]
            combo.column_indices = indices
            combo.set('')

        @staticmethod
        def _combo_col(combo):
            """Selected column index of a combobox filled by _set_column_choices."""
            cur = combo.current()
            indices = getattr(combo, 'column_indices', None)
            if cur is None or cur < 0 or not indices:
                return None
            return indices[cur]

        def _bind_mousewheel(self, area, canvas):
            """Cross-platform mouse wheel scrolling, active only while the pointer
            is over `area` (Windows/macOS <MouseWheel>, Linux Button-4/5)."""
            def on_wheel(event):
                if getattr(event, 'num', None) == 4:
                    delta = -1
                elif getattr(event, 'num', None) == 5:
                    delta = 1
                elif abs(event.delta) >= 120:
                    delta = -int(event.delta / 120)
                else:
                    delta = -int(event.delta) or (-1 if event.delta > 0 else 1)
                canvas.yview_scroll(delta, 'units')
                return 'break'

            def bind_wheel(_event):
                canvas.bind_all('<MouseWheel>', on_wheel)
                canvas.bind_all('<Button-4>', on_wheel)
                canvas.bind_all('<Button-5>', on_wheel)

            def unbind_wheel(_event):
                canvas.unbind_all('<MouseWheel>')
                canvas.unbind_all('<Button-4>')
                canvas.unbind_all('<Button-5>')

            area.bind('<Enter>', bind_wheel)
            area.bind('<Leave>', unbind_wheel)

        def _show_page_figure(self, page_key, fig, parent):
            """Display a figure on a page, closing/removing the previous one."""
            old = self.page_figures.pop(page_key, None)
            if old is not None:
                try:
                    old.canvas.get_tk_widget().destroy()
                except Exception:
                    pass
            canvas = FigureCanvasTkAgg(fig, parent)
            canvas.draw()
            canvas.get_tk_widget().pack(fill='both', expand=True, pady=10)
            self.page_figures[page_key] = fig
            return canvas

        # ---------- UI construction ----------
        def _build_ui(self):
            # --- Sidebar ---
            self.sidebar = ttk.Frame(self, style='TFrame', width=220)
            self.sidebar.pack(side='left', fill='y')
            self.sidebar.pack_propagate(False)

            logo_frame = ttk.Frame(self.sidebar)
            logo_frame.pack(fill='x', padx=15, pady=(20, 10))
            ttk.Label(logo_frame, text="📊", font=('Segoe UI', 28)).pack()
            ttk.Label(logo_frame, text="Accounting", style='Title.TLabel',
                      font=('Segoe UI', 14, 'bold')).pack()
            ttk.Label(logo_frame, text="Analyzer", style='Subtitle.TLabel',
                      font=('Segoe UI', 9)).pack()

            sep = tk.Frame(self.sidebar, height=2, bg=COLORS['border'])
            sep.pack(fill='x', padx=15, pady=10)

            self.nav_buttons = {}
            nav_items = [
                ('home', '🏠  Home'),
                ('data', '📋  Data'),
                ('financial', '💰  Financial Analysis'),
                ('categories', '📊  Breakdown'),
                ('reconciliation', '🔄  Reconciliation'),
                ('anomalies', '⚠️  Anomalies'),
                ('macros', '🧬  Macros'),
                ('charts', '📈  Charts'),
                ('export', '📄  Export PDF'),
            ]
            for key, label in nav_items:
                btn = ttk.Button(self.sidebar, text=label, style='Nav.TButton',
                                 command=lambda k=key: self.show_page(k))
                btn.pack(fill='x', padx=10, pady=2)
                self.nav_buttons[key] = btn

            self.autosave_label = ttk.Label(self.sidebar, text="",
                                            style='Subtitle.TLabel',
                                            font=('Segoe UI', 8))
            self.autosave_label.pack(side='bottom', padx=10, pady=(0, 8))
            self.file_info_label = ttk.Label(self.sidebar, text="No file loaded",
                                             style='Subtitle.TLabel', wraplength=190,
                                             font=('Segoe UI', 8))
            self.file_info_label.pack(side='bottom', padx=10, pady=15)

            # --- Main area ---
            self.main_area = ttk.Frame(self, style='TFrame')
            self.main_area.pack(side='right', fill='both', expand=True)

            self.pages = {}
            self.current_page = None

            self._build_home_page()
            self._build_data_page()
            self._build_financial_page()
            self._build_categories_page()
            self._build_reconciliation_page()
            self._build_anomalies_page()
            self._build_macros_page()
            self._build_charts_page()
            self._build_export_page()

            self.show_page('home')

        def _create_page(self, name):
            frame = ttk.Frame(self.main_area, style='TFrame')
            self.pages[name] = frame
            return frame

        def show_page(self, name):
            if self.current_page:
                self.pages[self.current_page].pack_forget()
            self.pages[name].pack(fill='both', expand=True)
            self.current_page = name

            for key, btn in self.nav_buttons.items():
                btn.configure(style='Accent.TButton' if key == name else 'Nav.TButton')

        # ---------- HOME PAGE ----------
        def _build_home_page(self):
            page = self._create_page('home')

            center = ttk.Frame(page, style='TFrame')
            center.place(relx=0.5, rely=0.4, anchor='center')

            ttk.Label(center, text="📊", font=('Segoe UI', 60)).pack(pady=(0, 10))
            ttk.Label(center, text="AccountingAnalyzer", style='Title.TLabel',
                      font=('Segoe UI', 28, 'bold')).pack()
            ttk.Label(center, text="Analyze your accounting spreadsheets in a few clicks",
                      style='Subtitle.TLabel', font=('Segoe UI', 12)).pack(pady=(5, 30))

            btn_frame = ttk.Frame(center, style='TFrame')
            btn_frame.pack()

            ttk.Button(btn_frame, text="📂  Open a Spreadsheet",
                       style='Accent.TButton', command=self.open_file).pack(pady=5)

            ttk.Label(center,
                      text="Supported formats: .xlsx  .xlsm  .xls  .xml  .csv  .tsv  .ods",
                      style='Subtitle.TLabel', font=('Segoe UI', 9)).pack(pady=(20, 0))
            ttk.Label(center,
                      text="Macro-enabled workbooks are analyzed statically — macros never run.",
                      style='Subtitle.TLabel', font=('Segoe UI', 8)).pack(pady=(4, 0))

        # ---------- DATA PAGE ----------
        def _build_data_page(self):
            page = self._create_page('data')

            header = ttk.Frame(page, style='TFrame')
            header.pack(fill='x', padx=20, pady=(15, 5))
            ttk.Label(header, text="📋 Raw Data", style='Section.TLabel').pack(side='left')

            ctrl = ttk.Frame(page, style='TFrame')
            ctrl.pack(fill='x', padx=20, pady=5)
            ttk.Label(ctrl, text="Sheet:").pack(side='left')
            self.sheet_combo = ttk.Combobox(ctrl, state='readonly', width=30)
            self.sheet_combo.pack(side='left', padx=10)
            self.sheet_combo.bind('<<ComboboxSelected>>', self._on_sheet_selected)

            self.data_info_label = ttk.Label(ctrl, text="", style='Subtitle.TLabel')
            self.data_info_label.pack(side='right')

            tree_frame = ttk.Frame(page)
            tree_frame.pack(fill='both', expand=True, padx=20, pady=10)

            self.data_tree = ttk.Treeview(tree_frame, show='headings')
            vsb = ttk.Scrollbar(tree_frame, orient='vertical', command=self.data_tree.yview)
            hsb = ttk.Scrollbar(tree_frame, orient='horizontal', command=self.data_tree.xview)
            self.data_tree.configure(yscrollcommand=vsb.set, xscrollcommand=hsb.set)

            self.data_tree.grid(row=0, column=0, sticky='nsew')
            vsb.grid(row=0, column=1, sticky='ns')
            hsb.grid(row=1, column=0, sticky='ew')
            tree_frame.grid_rowconfigure(0, weight=1)
            tree_frame.grid_columnconfigure(0, weight=1)

        def _on_sheet_selected(self, event=None):
            sheet_name = self.sheet_combo.get()
            if not sheet_name:
                return
            self.current_sheet = sheet_name
            self._populate_data_tree(sheet_name)
            self._save_session()

        MAX_GRID_ROWS = 500

        def _populate_data_tree(self, sheet_name):
            sheet = self.analyzer.get_sheet_data(sheet_name)
            if not sheet:
                return

            self.data_tree.delete(*self.data_tree.get_children())
            self.data_tree['columns'] = []

            headers = sheet['headers']
            if not headers:
                self.data_info_label.configure(text="Empty sheet")
                return

            cols = [f"col_{i}" for i in range(len(headers))]
            self.data_tree['columns'] = cols

            labels = self._column_labels(headers)
            for col_id, label in zip(cols, labels):
                self.data_tree.heading(col_id, text=label)
                self.data_tree.column(col_id, width=120, minwidth=80)

            for row in sheet['rows'][:self.MAX_GRID_ROWS]:
                values = [str(v) if v is not None else '' for v in row]
                while len(values) < len(cols):
                    values.append('')
                self.data_tree.insert('', 'end', values=values[:len(cols)])

            shown = min(sheet['num_rows'], self.MAX_GRID_ROWS)
            if sheet['num_rows'] > shown:
                text = (f"showing first {shown} of {sheet['num_rows']} rows "
                        f"× {sheet['num_cols']} columns")
            else:
                text = f"{sheet['num_rows']} rows × {sheet['num_cols']} columns"
            self.data_info_label.configure(text=text)

        # ---------- FINANCIAL ANALYSIS PAGE ----------
        def _build_financial_page(self):
            page = self._create_page('financial')

            header = ttk.Frame(page, style='TFrame')
            header.pack(fill='x', padx=20, pady=(15, 5))
            ttk.Label(header, text="💰 Financial Analysis", style='Section.TLabel').pack(side='left')

            ctrl = ttk.Frame(page, style='TFrame')
            ctrl.pack(fill='x', padx=20, pady=5)
            ttk.Label(ctrl, text="Sheet:").pack(side='left')
            self.fin_sheet_combo = ttk.Combobox(ctrl, state='readonly', width=30)
            self.fin_sheet_combo.pack(side='left', padx=10)
            ttk.Button(ctrl, text="Analyze", style='Secondary.TButton',
                       command=self._run_financial_analysis).pack(side='left', padx=5)

            canvas = tk.Canvas(page, bg=COLORS['bg'], highlightthickness=0)
            scrollbar = ttk.Scrollbar(page, orient='vertical', command=canvas.yview)
            self.fin_results_frame = ttk.Frame(canvas, style='TFrame')

            self.fin_results_frame.bind(
                '<Configure>', lambda e: canvas.configure(scrollregion=canvas.bbox('all'))
            )
            canvas.create_window((0, 0), window=self.fin_results_frame, anchor='nw')
            canvas.configure(yscrollcommand=scrollbar.set)

            canvas.pack(side='left', fill='both', expand=True, padx=20, pady=10)
            scrollbar.pack(side='right', fill='y')

            self._bind_mousewheel(page, canvas)

        def _run_financial_analysis(self):
            sheet_name = self.fin_sheet_combo.get()
            if not sheet_name:
                messagebox.showwarning("Warning", "Please select a sheet.")
                return

            for w in self.fin_results_frame.winfo_children():
                w.destroy()

            summary = self.analyzer.analyze_financial_summary(sheet_name)
            ratios = self.analyzer.compute_ratios(summary)
            self._save_session()

            if not summary:
                ttk.Label(self.fin_results_frame,
                          text="No numeric columns detected.",
                          style='Subtitle.TLabel').pack(pady=20)
                return

            kpi_frame = ttk.Frame(self.fin_results_frame, style='TFrame')
            kpi_frame.pack(fill='x', pady=(0, 15))

            col_count = 0
            for col_name, stats in summary.items():
                card = tk.Frame(kpi_frame, bg=COLORS['bg_card'], padx=15, pady=10)
                card.grid(row=col_count // 4, column=col_count % 4, padx=5, pady=5, sticky='nsew')
                kpi_frame.grid_columnconfigure(col_count % 4, weight=1)

                tk.Label(card, text=str(col_name)[:20], bg=COLORS['bg_card'],
                         fg=COLORS['text_dim'], font=('Segoe UI', 9)).pack()
                tk.Label(card, text=fmt_money(stats['total']), bg=COLORS['bg_card'],
                         fg=COLORS['accent'], font=('Segoe UI', 16, 'bold')).pack()
                tk.Label(card, text=f"Avg: {fmt_money(stats['average'])}",
                         bg=COLORS['bg_card'], fg=COLORS['text_dim'],
                         font=('Segoe UI', 8)).pack()
                col_count += 1

            ttk.Label(self.fin_results_frame, text="Detail by Column",
                      style='Section.TLabel', font=('Segoe UI', 12, 'bold')).pack(anchor='w', pady=(15, 5))

            tree = ttk.Treeview(self.fin_results_frame, show='headings',
                                height=min(len(summary), 10))
            cols = ('column', 'total', 'average', 'min', 'max', 'count', 'positives', 'negatives')
            tree['columns'] = cols
            for c in cols:
                tree.heading(c, text=c.capitalize())
                tree.column(c, width=110)

            for col_name, stats in summary.items():
                tree.insert('', 'end', values=(
                    str(col_name)[:25],
                    fmt_money(stats['total']),
                    fmt_money(stats['average']),
                    fmt_money(stats['min']),
                    fmt_money(stats['max']),
                    stats['count'],
                    fmt_money(stats['total_positives']),
                    fmt_money(stats['total_negatives'])
                ))
            tree.pack(fill='x', pady=5)

            if ratios:
                ttk.Label(self.fin_results_frame, text="Computed Ratios",
                          style='Section.TLabel', font=('Segoe UI', 12, 'bold')).pack(anchor='w', pady=(15, 5))

                ratio_frame = ttk.Frame(self.fin_results_frame, style='TFrame')
                ratio_frame.pack(fill='x')
                for i, (name, val) in enumerate(ratios.items()):
                    card = tk.Frame(ratio_frame, bg=COLORS['bg_card'], padx=12, pady=8)
                    card.grid(row=i // 3, column=i % 3, padx=5, pady=3, sticky='nsew')
                    ratio_frame.grid_columnconfigure(i % 3, weight=1)
                    tk.Label(card, text=name[:30], bg=COLORS['bg_card'],
                             fg=COLORS['text_dim'], font=('Segoe UI', 8)).pack()
                    color = COLORS['success'] if val >= 0 else COLORS['danger']
                    tk.Label(card,
                             text=fmt_pct(val) if 'ratio' in name.lower() or '%' in name else fmt_money(val),
                             bg=COLORS['bg_card'], fg=color,
                             font=('Segoe UI', 14, 'bold')).pack()

        # ---------- BREAKDOWN PAGE ----------
        def _build_categories_page(self):
            page = self._create_page('categories')

            header = ttk.Frame(page, style='TFrame')
            header.pack(fill='x', padx=20, pady=(15, 5))
            ttk.Label(header, text="📊 Breakdown by Category", style='Section.TLabel').pack(side='left')

            ctrl = ttk.Frame(page, style='TFrame')
            ctrl.pack(fill='x', padx=20, pady=5)

            ttk.Label(ctrl, text="Sheet:").pack(side='left')
            self.cat_sheet_combo = ttk.Combobox(ctrl, state='readonly', width=20)
            self.cat_sheet_combo.pack(side='left', padx=5)
            self.cat_sheet_combo.bind('<<ComboboxSelected>>', self._update_cat_cols)

            ttk.Label(ctrl, text="Category:").pack(side='left', padx=(10, 0))
            self.cat_col_combo = ttk.Combobox(ctrl, state='readonly', width=20)
            self.cat_col_combo.pack(side='left', padx=5)

            ttk.Label(ctrl, text="Value:").pack(side='left', padx=(10, 0))
            self.cat_val_combo = ttk.Combobox(ctrl, state='readonly', width=20)
            self.cat_val_combo.pack(side='left', padx=5)

            ttk.Button(ctrl, text="Analyze", style='Secondary.TButton',
                       command=self._run_category_analysis).pack(side='left', padx=10)

            self.cat_results_frame = ttk.Frame(page, style='TFrame')
            self.cat_results_frame.pack(fill='both', expand=True, padx=20, pady=10)

        def _update_cat_cols(self, event=None):
            sheet = self.cat_sheet_combo.get()
            if not sheet:
                return
            headers = self.analyzer.get_sheet_data(sheet).get('headers', [])
            self._set_column_choices(self.cat_col_combo, headers)
            self._set_column_choices(self.cat_val_combo, headers)

        def _run_category_analysis(self):
            sheet = self.cat_sheet_combo.get()
            cat_idx = self._combo_col(self.cat_col_combo)
            val_idx = self._combo_col(self.cat_val_combo)
            if not sheet or cat_idx is None or val_idx is None:
                messagebox.showwarning("Warning", "Please fill in all fields.")
                return

            results = self.analyzer.analyze_by_category(sheet, cat_idx, val_idx)
            self._save_session()

            for w in self.cat_results_frame.winfo_children():
                w.destroy()
            self.page_figures.pop('categories', None)

            if not results:
                ttk.Label(self.cat_results_frame, text="No results.",
                          style='Subtitle.TLabel').pack(pady=20)
                return

            tree = ttk.Treeview(self.cat_results_frame, show='headings',
                                height=min(len(results), 15))
            cols = ('category', 'total', 'average', 'min', 'max', 'count', 'share')
            tree['columns'] = cols
            labels = ('Category', 'Total', 'Average', 'Min', 'Max', 'Count', 'Share %')
            for c, l in zip(cols, labels):
                tree.heading(c, text=l)
                tree.column(c, width=110)

            grand_total = sum(abs(v['total']) for v in results.values())
            sorted_results = sorted(results.items(), key=lambda x: abs(x[1]['total']),
                                    reverse=True)

            for cat, stats in sorted_results:
                pct = (abs(stats['total']) / grand_total * 100) if grand_total else 0
                tree.insert('', 'end', values=(
                    str(cat)[:30], fmt_money(stats['total']), fmt_money(stats['average']),
                    fmt_money(stats['min']), fmt_money(stats['max']),
                    stats['count'], fmt_pct(pct)
                ))
            tree.pack(fill='x', pady=5)

            fig = Figure(figsize=(8, 4), dpi=100, facecolor=COLORS['bg'])
            ax = fig.add_subplot(121)
            ax2 = fig.add_subplot(122)

            top_n = sorted_results[:8]
            labels_pie = [str(k)[:15] for k, _ in top_n]
            values_pie = [abs(v['total']) for _, v in top_n]
            if len(sorted_results) > 8:
                labels_pie.append('Others')
                values_pie.append(sum(abs(v['total']) for _, v in sorted_results[8:]))

            ax.pie(values_pie, labels=labels_pie, autopct='%1.1f%%',
                   colors=COLORS['chart_colors'][:len(labels_pie)],
                   textprops={'color': COLORS['text'], 'fontsize': 8})
            ax.set_title('Distribution', color=COLORS['text'], fontsize=11)
            ax.set_facecolor(COLORS['bg'])

            bar_labels = [str(k)[:12] for k, _ in top_n]
            bar_values = [v['total'] for _, v in top_n]
            bar_colors = [COLORS['success'] if v >= 0 else COLORS['danger'] for v in bar_values]
            ax2.barh(bar_labels[::-1], bar_values[::-1], color=bar_colors[::-1])
            ax2.set_title('Top Categories', color=COLORS['text'], fontsize=11)
            ax2.set_facecolor(COLORS['bg_light'])
            ax2.tick_params(colors=COLORS['text'])
            for spine in ax2.spines.values():
                spine.set_color(COLORS['border'])

            fig.tight_layout()
            self._show_page_figure('categories', fig, self.cat_results_frame)

        # ---------- RECONCILIATION PAGE ----------
        def _build_reconciliation_page(self):
            page = self._create_page('reconciliation')

            header = ttk.Frame(page, style='TFrame')
            header.pack(fill='x', padx=20, pady=(15, 5))
            ttk.Label(header, text="🔄 Bank Reconciliation", style='Section.TLabel').pack(side='left')

            ctrl = ttk.Frame(page, style='TFrame')
            ctrl.pack(fill='x', padx=20, pady=5)

            f1 = ttk.Frame(ctrl, style='TFrame')
            f1.pack(fill='x', pady=3)
            ttk.Label(f1, text="Sheet 1:").pack(side='left')
            self.rec_sheet1_combo = ttk.Combobox(f1, state='readonly', width=20)
            self.rec_sheet1_combo.pack(side='left', padx=5)
            self.rec_sheet1_combo.bind('<<ComboboxSelected>>', lambda e: self._update_rec_cols(1))
            ttk.Label(f1, text="Column:").pack(side='left', padx=(10, 0))
            self.rec_col1_combo = ttk.Combobox(f1, state='readonly', width=20)
            self.rec_col1_combo.pack(side='left', padx=5)

            f2 = ttk.Frame(ctrl, style='TFrame')
            f2.pack(fill='x', pady=3)
            ttk.Label(f2, text="Sheet 2:").pack(side='left')
            self.rec_sheet2_combo = ttk.Combobox(f2, state='readonly', width=20)
            self.rec_sheet2_combo.pack(side='left', padx=5)
            self.rec_sheet2_combo.bind('<<ComboboxSelected>>', lambda e: self._update_rec_cols(2))
            ttk.Label(f2, text="Column:").pack(side='left', padx=(10, 0))
            self.rec_col2_combo = ttk.Combobox(f2, state='readonly', width=20)
            self.rec_col2_combo.pack(side='left', padx=5)

            ttk.Button(ctrl, text="Run Reconciliation", style='Secondary.TButton',
                       command=self._run_reconciliation).pack(pady=10)

            self.rec_results_frame = ttk.Frame(page, style='TFrame')
            self.rec_results_frame.pack(fill='both', expand=True, padx=20, pady=10)

        def _update_rec_cols(self, num):
            combo = self.rec_sheet1_combo if num == 1 else self.rec_sheet2_combo
            col_combo = self.rec_col1_combo if num == 1 else self.rec_col2_combo
            sheet = combo.get()
            if sheet:
                headers = self.analyzer.get_sheet_data(sheet).get('headers', [])
                self._set_column_choices(col_combo, headers)

        def _run_reconciliation(self):
            s1 = self.rec_sheet1_combo.get()
            s2 = self.rec_sheet2_combo.get()
            idx1 = self._combo_col(self.rec_col1_combo)
            idx2 = self._combo_col(self.rec_col2_combo)

            if not s1 or not s2 or idx1 is None or idx2 is None:
                messagebox.showwarning("Warning", "Please fill in all fields.")
                return

            result = self.analyzer.bank_reconciliation(s1, idx1, s2, idx2)
            self._save_session()

            for w in self.rec_results_frame.winfo_children():
                w.destroy()

            kpi = ttk.Frame(self.rec_results_frame, style='TFrame')
            kpi.pack(fill='x', pady=10)

            kpis = [
                ("Matched", str(len(result['matched'])), COLORS['success']),
                ("Unmatched S1", str(len(result['unmatched_sheet1'])), COLORS['warning']),
                ("Unmatched S2", str(len(result['unmatched_sheet2'])), COLORS['warning']),
                ("Total S1", fmt_money(result['total_sheet1']), COLORS['text']),
                ("Total S2", fmt_money(result['total_sheet2']), COLORS['text']),
                ("Difference", fmt_money(result['difference']),
                 COLORS['danger'] if abs(result['difference']) > 0.01 else COLORS['success']),
                ("Match Rate", fmt_pct(result['match_rate']),
                 COLORS['success'] if result['match_rate'] > 90 else COLORS['danger']),
            ]

            for i, (label, value, color) in enumerate(kpis):
                card = tk.Frame(kpi, bg=COLORS['bg_card'], padx=12, pady=8)
                card.grid(row=0, column=i, padx=4, pady=4, sticky='nsew')
                kpi.grid_columnconfigure(i, weight=1)
                tk.Label(card, text=label, bg=COLORS['bg_card'],
                         fg=COLORS['text_dim'], font=('Segoe UI', 8)).pack()
                tk.Label(card, text=value, bg=COLORS['bg_card'],
                         fg=color, font=('Segoe UI', 13, 'bold')).pack()

            for key, title in (('unmatched_sheet1', f"Unmatched — {s1}"),
                               ('unmatched_sheet2', f"Unmatched — {s2}")):
                if result[key]:
                    ttk.Label(self.rec_results_frame, text=title,
                              style='Section.TLabel',
                              font=('Segoe UI', 11, 'bold')).pack(anchor='w', pady=(10, 3))
                    tree = ttk.Treeview(self.rec_results_frame, show='headings',
                                        height=min(len(result[key]), 8))
                    tree['columns'] = ('row', 'value')
                    tree.heading('row', text='Row')
                    tree.heading('value', text='Amount')
                    tree.column('row', width=100)
                    tree.column('value', width=200)
                    for item in result[key]:
                        tree.insert('', 'end', values=(item['row'], fmt_money(item['value'])))
                    tree.pack(fill='x', pady=3)

        # ---------- ANOMALIES PAGE ----------
        def _build_anomalies_page(self):
            page = self._create_page('anomalies')

            header = ttk.Frame(page, style='TFrame')
            header.pack(fill='x', padx=20, pady=(15, 5))
            ttk.Label(header, text="⚠️ Anomaly Detection", style='Section.TLabel').pack(side='left')

            ctrl = ttk.Frame(page, style='TFrame')
            ctrl.pack(fill='x', padx=20, pady=5)
            ttk.Label(ctrl, text="Sheet:").pack(side='left')
            self.anom_sheet_combo = ttk.Combobox(ctrl, state='readonly', width=20)
            self.anom_sheet_combo.pack(side='left', padx=5)
            self.anom_sheet_combo.bind('<<ComboboxSelected>>', self._update_anom_cols)
            ttk.Label(ctrl, text="Column:").pack(side='left', padx=(10, 0))
            self.anom_col_combo = ttk.Combobox(ctrl, state='readonly', width=20)
            self.anom_col_combo.pack(side='left', padx=5)
            ttk.Label(ctrl, text="Threshold (σ):").pack(side='left', padx=(10, 0))
            self.anom_threshold = ttk.Combobox(ctrl, state='readonly', width=5,
                                               values=['1.5', '2.0', '2.5', '3.0'])
            self.anom_threshold.set('2.0')
            self.anom_threshold.pack(side='left', padx=5)
            ttk.Button(ctrl, text="Detect", style='Secondary.TButton',
                       command=self._run_anomaly_detection).pack(side='left', padx=10)

            self.anom_results_frame = ttk.Frame(page, style='TFrame')
            self.anom_results_frame.pack(fill='both', expand=True, padx=20, pady=10)

        def _update_anom_cols(self, event=None):
            sheet = self.anom_sheet_combo.get()
            if sheet:
                headers = self.analyzer.get_sheet_data(sheet).get('headers', [])
                numeric_indices = [i for i, _ in self.analyzer.detect_numeric_columns(sheet)]
                self._set_column_choices(self.anom_col_combo, headers, numeric_indices)

        def _run_anomaly_detection(self):
            sheet = self.anom_sheet_combo.get()
            col_idx = self._combo_col(self.anom_col_combo)
            threshold = float(self.anom_threshold.get() or 2.0)

            if not sheet or col_idx is None:
                messagebox.showwarning("Warning", "Please fill in all fields.")
                return

            anomalies = self.analyzer.detect_anomalies(sheet, col_idx, threshold)
            self._save_session()

            for w in self.anom_results_frame.winfo_children():
                w.destroy()

            if not anomalies:
                ttk.Label(self.anom_results_frame,
                          text=f"✅ No anomalies detected (threshold = {threshold}σ)",
                          style='Subtitle.TLabel', font=('Segoe UI', 12)).pack(pady=30)
                return

            ttk.Label(self.anom_results_frame,
                      text=f"⚠️ {len(anomalies)} anomaly(ies) detected",
                      foreground=COLORS['warning'],
                      font=('Segoe UI', 13, 'bold')).pack(anchor='w', pady=(0, 10))

            tree = ttk.Treeview(self.anom_results_frame, show='headings',
                                height=min(len(anomalies), 15))
            tree['columns'] = ('row', 'value', 'mean', 'deviation', 'zscore')
            tree.heading('row', text='Row')
            tree.heading('value', text='Value')
            tree.heading('mean', text='Mean')
            tree.heading('deviation', text='Deviation')
            tree.heading('zscore', text='Z-Score')
            for c in tree['columns']:
                tree.column(c, width=130)

            for a in anomalies:
                tree.insert('', 'end', values=(
                    a['row'], fmt_money(a['value']), fmt_money(a['mean']),
                    fmt_money(a['deviation']), f"{a['z_score']:.2f}"
                ))
            tree.pack(fill='x')

        # ---------- MACROS PAGE ----------
        def _build_macros_page(self):
            page = self._create_page('macros')

            header = ttk.Frame(page, style='TFrame')
            header.pack(fill='x', padx=20, pady=(15, 5))
            ttk.Label(header, text="🧬 Macro Analysis (static — macros never run)",
                      style='Section.TLabel').pack(side='left')

            self.macro_text = tk.Text(page, bg=COLORS['bg_light'], fg=COLORS['text'],
                                      insertbackground=COLORS['text'],
                                      font=('Consolas', 10), wrap='word',
                                      relief='flat', padx=12, pady=12)
            self.macro_text.pack(fill='both', expand=True, padx=20, pady=10)
            self.macro_text.configure(state='disabled')
            self._refresh_macro_page()

        def _refresh_macro_page(self):
            info = self.analyzer.macro_info
            lines = [info.get('summary') or 'No file loaded.', '']
            if info.get('error'):
                lines += [f"Note: {info['error']}", '']
            if info.get('modules'):
                lines.append('Modules:')
                lines += [f"  • {m}" for m in info['modules']]
                lines.append('')
            if info.get('findings'):
                lines.append('Findings (from oletools static scan):')
                for f in info['findings']:
                    lines.append(f"  [{f['type']}] {f['keyword']} — {f['description']}")
            self.macro_text.configure(state='normal')
            self.macro_text.delete('1.0', 'end')
            self.macro_text.insert('1.0', '\n'.join(lines))
            self.macro_text.configure(state='disabled')

        # ---------- CHARTS PAGE ----------
        def _build_charts_page(self):
            page = self._create_page('charts')

            header = ttk.Frame(page, style='TFrame')
            header.pack(fill='x', padx=20, pady=(15, 5))
            ttk.Label(header, text="📈 Charts", style='Section.TLabel').pack(side='left')

            ctrl = ttk.Frame(page, style='TFrame')
            ctrl.pack(fill='x', padx=20, pady=5)

            ttk.Label(ctrl, text="Sheet:").pack(side='left')
            self.chart_sheet_combo = ttk.Combobox(ctrl, state='readonly', width=20)
            self.chart_sheet_combo.pack(side='left', padx=5)
            self.chart_sheet_combo.bind('<<ComboboxSelected>>', self._update_chart_cols)

            ttk.Label(ctrl, text="Type:").pack(side='left', padx=(10, 0))
            self.chart_type_combo = ttk.Combobox(ctrl, state='readonly', width=15,
                                                 values=['Bar', 'Line', 'Pie',
                                                         'Area', 'Histogram', 'Scatter'])
            self.chart_type_combo.set('Bar')
            self.chart_type_combo.pack(side='left', padx=5)

            ttk.Label(ctrl, text="X Axis:").pack(side='left', padx=(10, 0))
            self.chart_x_combo = ttk.Combobox(ctrl, state='readonly', width=18)
            self.chart_x_combo.pack(side='left', padx=5)

            ttk.Label(ctrl, text="Y Axis:").pack(side='left', padx=(10, 0))
            self.chart_y_combo = ttk.Combobox(ctrl, state='readonly', width=18)
            self.chart_y_combo.pack(side='left', padx=5)

            ttk.Button(ctrl, text="Generate", style='Secondary.TButton',
                       command=self._generate_chart).pack(side='left', padx=10)

            self.chart_info_label = ttk.Label(page, text="", style='Subtitle.TLabel')
            self.chart_info_label.pack(anchor='w', padx=20)

            self.chart_canvas_frame = ttk.Frame(page, style='TFrame')
            self.chart_canvas_frame.pack(fill='both', expand=True, padx=20, pady=10)

        def _update_chart_cols(self, event=None):
            sheet = self.chart_sheet_combo.get()
            if sheet:
                headers = self.analyzer.get_sheet_data(sheet).get('headers', [])
                self._set_column_choices(self.chart_x_combo, headers)
                self._set_column_choices(self.chart_y_combo, headers)

        MAX_CHART_ITEMS = 30

        def _generate_chart(self):
            sheet = self.chart_sheet_combo.get()
            chart_type = self.chart_type_combo.get()
            x_idx = self._combo_col(self.chart_x_combo)
            y_idx = self._combo_col(self.chart_y_combo)

            if not sheet or not chart_type or x_idx is None or y_idx is None:
                messagebox.showwarning("Warning", "Please fill in all fields.")
                return

            data = self.analyzer.get_sheet_data(sheet)
            self._save_session()

            date_col_indices = {i for i, _ in self.analyzer.detect_date_columns(sheet)}
            use_dates = chart_type in ('Line', 'Area') and x_idx in date_col_indices

            if use_dates:
                series = self.analyzer.analyze_time_series(sheet, x_idx, y_idx)
                x_vals = [d.strftime('%m/%d/%Y') for d, _ in series]
                y_vals = [v for _, v in series]
            else:
                x_vals, y_vals = [], []
                for row in data.get('rows', []):
                    if x_idx < len(row) and y_idx < len(row):
                        x = row[x_idx]
                        y = row[y_idx]
                        if x is not None and y is not None:
                            try:
                                y_val = to_number(y)
                            except (ValueError, TypeError):
                                continue
                            x_vals.append(x if is_number(x) else str(x))
                            y_vals.append(y_val)

            if not x_vals:
                messagebox.showwarning("Warning", "No data to display.")
                return

            for w in self.chart_canvas_frame.winfo_children():
                w.destroy()
            self.page_figures.pop('charts', None)

            truncated = False
            if len(x_vals) > self.MAX_CHART_ITEMS and chart_type in ['Bar', 'Pie']:
                x_vals = x_vals[:self.MAX_CHART_ITEMS]
                y_vals = y_vals[:self.MAX_CHART_ITEMS]
                truncated = True

            fig = Figure(figsize=(10, 5), dpi=100, facecolor=COLORS['bg'])
            ax = fig.add_subplot(111)
            ax.set_facecolor(COLORS['bg_light'])

            if chart_type == 'Bar':
                bar_colors = [COLORS['success'] if v >= 0 else COLORS['danger'] for v in y_vals]
                ax.bar(range(len(x_vals)), y_vals, color=bar_colors)
                ax.set_xticks(range(len(x_vals)))
                ax.set_xticklabels([str(x)[:10] for x in x_vals], rotation=45, ha='right',
                                   fontsize=7, color=COLORS['text'])
            elif chart_type == 'Line':
                ax.plot(range(len(x_vals)), y_vals, color=COLORS['accent'],
                        linewidth=2, marker='o', markersize=4)
                ax.fill_between(range(len(y_vals)), y_vals, alpha=0.1, color=COLORS['accent'])
                step = max(1, len(x_vals) // 10)
                ax.set_xticks(range(0, len(x_vals), step))
                ax.set_xticklabels([str(x_vals[i])[:10] for i in range(0, len(x_vals), step)],
                                   rotation=45, fontsize=7, color=COLORS['text'])
            elif chart_type == 'Pie':
                ax.pie(y_vals, labels=[str(x)[:12] for x in x_vals],
                       autopct='%1.1f%%', colors=COLORS['chart_colors'][:len(x_vals)],
                       textprops={'color': COLORS['text'], 'fontsize': 8})
            elif chart_type == 'Area':
                ax.fill_between(range(len(y_vals)), y_vals, alpha=0.5, color=COLORS['accent'])
                ax.plot(range(len(y_vals)), y_vals, color=COLORS['accent'], linewidth=1)
            elif chart_type == 'Histogram':
                ax.hist(y_vals, bins=min(20, len(y_vals)), color=COLORS['accent'],
                        edgecolor=COLORS['bg'], alpha=0.8)
            elif chart_type == 'Scatter':
                try:
                    x_numeric = [float(x) for x in x_vals]
                    ax.scatter(x_numeric, y_vals, color=COLORS['accent'], alpha=0.7, s=30)
                except (ValueError, TypeError):
                    ax.scatter(range(len(y_vals)), y_vals, color=COLORS['accent'], alpha=0.7, s=30)

            x_name = self.chart_x_combo.get()
            y_name = self.chart_y_combo.get()
            ax.set_title(f"{y_name} by {x_name}", color=COLORS['text'], fontsize=12, pad=10)
            ax.tick_params(colors=COLORS['text'])
            for spine in ax.spines.values():
                spine.set_color(COLORS['border'])

            fig.tight_layout()
            self._show_page_figure('charts', fig, self.chart_canvas_frame)

            info = []
            if truncated:
                info.append(f"showing first {self.MAX_CHART_ITEMS} rows")
            if use_dates:
                info.append("sorted chronologically")
            self.chart_info_label.configure(text=' — '.join(info))

        # ---------- EXPORT PDF PAGE ----------
        def _build_export_page(self):
            page = self._create_page('export')

            header = ttk.Frame(page, style='TFrame')
            header.pack(fill='x', padx=20, pady=(15, 5))
            ttk.Label(header, text="📄 Export PDF", style='Section.TLabel').pack(side='left')

            center = ttk.Frame(page, style='TFrame')
            center.place(relx=0.5, rely=0.4, anchor='center')

            ttk.Label(center, text="📄", font=('Segoe UI', 50)).pack()
            ttk.Label(center, text="Generate a Complete PDF Report",
                      style='Title.TLabel', font=('Segoe UI', 16, 'bold')).pack(pady=(10, 5))
            ttk.Label(center,
                      text="The report can include: financial summary, breakdown,\n"
                           "detected anomalies, charts and macro analysis",
                      style='Subtitle.TLabel', justify='center').pack(pady=(0, 20))

            options_frame = ttk.Frame(center, style='TFrame')
            options_frame.pack(pady=10)

            self.export_vars = {}
            for text, key in [("Financial Summary", 'summary'),
                              ("Category Breakdown", 'categories'),
                              ("Anomalies", 'anomalies'),
                              ("Charts", 'charts'),
                              ("Macro Analysis", 'macros')]:
                var = tk.BooleanVar(value=True)
                self.export_vars[key] = var
                cb = tk.Checkbutton(options_frame, text=text, variable=var,
                                    bg=COLORS['bg'], fg=COLORS['text'],
                                    selectcolor=COLORS['bg_card'],
                                    activebackground=COLORS['bg'],
                                    activeforeground=COLORS['text'],
                                    font=('Segoe UI', 10))
                cb.pack(anchor='w', pady=2)

            ttk.Button(center, text="📄  Generate PDF Report",
                       style='Accent.TButton', command=self._export_pdf).pack(pady=20)

            self.export_status = ttk.Label(center, text="", style='Subtitle.TLabel')
            self.export_status.pack()

        def _export_pdf(self):
            if not self.analyzer.sheets:
                messagebox.showwarning("Warning", "Please load a spreadsheet first.")
                return

            path = filedialog.asksaveasfilename(
                defaultextension='.pdf',
                filetypes=[('PDF', '*.pdf')],
                initialfile=f"accounting_report_{datetime.now().strftime('%Y%m%d')}.pdf"
            )
            if not path:
                return

            self.export_status.configure(text="Generating report...")
            self.update_idletasks()

            try:
                include = {key for key, var in self.export_vars.items() if var.get()}
                threshold = float(self.anom_threshold.get() or 2.0)
                build_pdf_report(self.analyzer, path, include=include,
                                 anomaly_threshold=threshold)
                self.export_status.configure(text=f"✅ Report exported: {os.path.basename(path)}")
                messagebox.showinfo("Success", f"PDF report generated:\n{path}")
            except Exception as e:
                self.export_status.configure(text=f"❌ Error: {str(e)}")
                messagebox.showerror("Error", f"Error during generation:\n{str(e)}")

        # ---------- FILE OPENING ----------
        def open_file(self):
            path = filedialog.askopenfilename(
                title="Open a Spreadsheet",
                filetypes=self.FILE_TYPES
            )
            if not path:
                return
            self._load_file(path)

        def _load_file(self, path, show_dialogs=True):
            try:
                sheets = self.analyzer.load_file(path)
            except Exception as e:
                if show_dialogs:
                    messagebox.showerror("Error", f"Unable to load file:\n{str(e)}")
                    return False
                raise

            sheet_names = list(sheets.keys())

            for combo in [self.sheet_combo, self.fin_sheet_combo, self.cat_sheet_combo,
                          self.anom_sheet_combo, self.chart_sheet_combo,
                          self.rec_sheet1_combo, self.rec_sheet2_combo]:
                combo['values'] = sheet_names
                if sheet_names:
                    combo.set(sheet_names[0])

            if sheet_names:
                self.current_sheet = sheet_names[0]
                self._populate_data_tree(sheet_names[0])
                self._update_cat_cols()
                self._update_anom_cols()
                self._update_chart_cols()
                self._update_rec_cols(1)
                self._update_rec_cols(2)

            self._refresh_macro_page()

            total_rows = sum(s['num_rows'] for s in sheets.values())
            macro_line = "🧬 macros detected" if self.analyzer.macro_info['has_macros'] else ""
            self.file_info_label.configure(
                text=f"📁 {os.path.basename(path)}\n"
                     f"{len(sheet_names)} sheet(s) — {total_rows} rows\n{macro_line}".rstrip()
            )

            self.show_page('data')
            self._save_session()

            if show_dialogs:
                message = (f"{os.path.basename(path)}\n\n"
                           f"Sheets: {len(sheet_names)}\n"
                           f"Total rows: {total_rows}")
                if self.analyzer.warnings:
                    message += "\n\nWarnings:\n" + "\n".join(self.analyzer.warnings)
                if self.analyzer.macro_info['has_macros']:
                    message += f"\n\n{self.analyzer.macro_info['summary']}"
                messagebox.showinfo("File Loaded", message)
            return True

        # ---------- SESSION AUTOSAVE ----------
        def _collect_state(self):
            return {
                'file_path': self.analyzer.file_path,
                'current_page': self.current_page,
                'current_sheet': self.current_sheet,
                'geometry': self.geometry(),
                'selections': {
                    'fin_sheet': self.fin_sheet_combo.current(),
                    'cat_sheet': self.cat_sheet_combo.current(),
                    'cat_col': self.cat_col_combo.current(),
                    'cat_val': self.cat_val_combo.current(),
                    'anom_sheet': self.anom_sheet_combo.current(),
                    'anom_col': self.anom_col_combo.current(),
                    'anom_threshold': self.anom_threshold.get(),
                    'chart_sheet': self.chart_sheet_combo.current(),
                    'chart_type': self.chart_type_combo.get(),
                    'chart_x': self.chart_x_combo.current(),
                    'chart_y': self.chart_y_combo.current(),
                    'rec_sheet1': self.rec_sheet1_combo.current(),
                    'rec_col1': self.rec_col1_combo.current(),
                    'rec_sheet2': self.rec_sheet2_combo.current(),
                    'rec_col2': self.rec_col2_combo.current(),
                },
                'export': {key: var.get() for key, var in self.export_vars.items()},
            }

        def _save_session(self):
            try:
                saved_at = self.session.save(self._collect_state())
                self.autosave_label.configure(
                    text=f"💾 autosaved {saved_at.split('T')[1]}")
            except OSError:
                self.autosave_label.configure(text="💾 autosave failed")

        def _autosave_tick(self):
            if self.analyzer.file_path:
                self._save_session()
            self.after(60000, self._autosave_tick)

        def _offer_session_restore(self):
            state = self.session.load()
            if not state:
                return
            path = state.get('file_path')
            if not path or not os.path.isfile(path):
                return
            saved_at = state.get('saved_at', 'unknown time')
            if not messagebox.askyesno(
                    "Restore Session",
                    f"Restore your last session?\n\n"
                    f"File: {os.path.basename(path)}\n"
                    f"Saved: {saved_at}"):
                return
            if not self._load_file(path, show_dialogs=False):
                return
            self._apply_state(state)

        @staticmethod
        def _restore_combo(combo, index):
            values = combo['values']
            if isinstance(index, int) and 0 <= index < len(values):
                combo.current(index)

        def _apply_state(self, state):
            sel = state.get('selections', {})
            self._restore_combo(self.fin_sheet_combo, sel.get('fin_sheet'))
            self._restore_combo(self.cat_sheet_combo, sel.get('cat_sheet'))
            self._update_cat_cols()
            self._restore_combo(self.cat_col_combo, sel.get('cat_col'))
            self._restore_combo(self.cat_val_combo, sel.get('cat_val'))
            self._restore_combo(self.anom_sheet_combo, sel.get('anom_sheet'))
            self._update_anom_cols()
            self._restore_combo(self.anom_col_combo, sel.get('anom_col'))
            if sel.get('anom_threshold'):
                self.anom_threshold.set(sel['anom_threshold'])
            self._restore_combo(self.chart_sheet_combo, sel.get('chart_sheet'))
            self._update_chart_cols()
            if sel.get('chart_type'):
                self.chart_type_combo.set(sel['chart_type'])
            self._restore_combo(self.chart_x_combo, sel.get('chart_x'))
            self._restore_combo(self.chart_y_combo, sel.get('chart_y'))
            self._restore_combo(self.rec_sheet1_combo, sel.get('rec_sheet1'))
            self._update_rec_cols(1)
            self._restore_combo(self.rec_col1_combo, sel.get('rec_col1'))
            self._restore_combo(self.rec_sheet2_combo, sel.get('rec_sheet2'))
            self._update_rec_cols(2)
            self._restore_combo(self.rec_col2_combo, sel.get('rec_col2'))
            for key, value in state.get('export', {}).items():
                if key in self.export_vars:
                    self.export_vars[key].set(bool(value))
            sheet = state.get('current_sheet')
            if sheet and sheet in self.analyzer.sheets:
                self.current_sheet = sheet
                self.sheet_combo.set(sheet)
                self._populate_data_tree(sheet)
            geometry = state.get('geometry')
            if geometry:
                try:
                    self.geometry(geometry)
                except tk.TclError:
                    pass
            page = state.get('current_page')
            if page in self.pages:
                self.show_page(page)

        def _on_close(self):
            if self.analyzer.file_path:
                self._save_session()
            self.destroy()

    def run_smoke_test(path):
        """Scripted GUI run: load a file, exercise every analysis page, export a
        PDF, exit 0 on success. Requires a display (use xvfb-run on CI)."""
        import tempfile

        app = AccountingApp(restore_session=False)
        try:
            app.update()
            assert app._load_file(path, show_dialogs=False), "load failed"
            app.update()

            sheets = app.analyzer.get_sheet_names()
            assert sheets, "no sheets loaded"

            # Financial analysis
            app.fin_sheet_combo.current(0)
            app._run_financial_analysis()
            app.update()

            # Category breakdown (first detected category × first numeric column)
            first = sheets[0]
            cat_cols = app.analyzer.detect_category_columns(first)
            num_cols = app.analyzer.detect_numeric_columns(first)
            if cat_cols and num_cols:
                app.cat_sheet_combo.current(0)
                app._update_cat_cols()
                indices = app.cat_col_combo.column_indices
                app.cat_col_combo.current(indices.index(cat_cols[0][0]))
                app.cat_val_combo.current(indices.index(num_cols[0][0]))
                app._run_category_analysis()
                app.update()

            # Reconciliation (sheet 1 against last sheet)
            if num_cols:
                app.rec_sheet1_combo.current(0)
                app._update_rec_cols(1)
                app.rec_sheet2_combo.current(len(sheets) - 1)
                app._update_rec_cols(2)
                idx1 = app.rec_col1_combo.column_indices
                app.rec_col1_combo.current(idx1.index(num_cols[0][0]))
                s2 = app.rec_sheet2_combo.get()
                num2 = app.analyzer.detect_numeric_columns(s2)
                if num2:
                    idx2 = app.rec_col2_combo.column_indices
                    app.rec_col2_combo.current(idx2.index(num2[0][0]))
                    app._run_reconciliation()
                    app.update()

            # Anomalies
            if num_cols:
                app.anom_sheet_combo.current(0)
                app._update_anom_cols()
                if app.anom_col_combo['values']:
                    app.anom_col_combo.current(0)
                    app._run_anomaly_detection()
                    app.update()

            # Chart
            if num_cols:
                app.chart_sheet_combo.current(0)
                app._update_chart_cols()
                app.chart_x_combo.current(0)
                app.chart_y_combo.current(num_cols[0][0])
                app._generate_chart()
                app.update()

            # PDF export via the shared report builder
            with tempfile.TemporaryDirectory() as tmp:
                out = os.path.join(tmp, 'smoke_report.pdf')
                build_pdf_report(app.analyzer, out)
                assert os.path.getsize(out) > 1024, "PDF too small"

            print("SMOKE TEST PASSED")
            return 0
        finally:
            app.destroy()


# =====================================================
# ENTRY POINT
# =====================================================
def main(argv=None):
    parser = argparse.ArgumentParser(
        description=f"AccountingAnalyzer v{VERSION} — spreadsheet analysis for accountants")
    parser.add_argument('--self-test', metavar='PATH',
                        help='headless: batch-load every spreadsheet under PATH and report')
    parser.add_argument('--smoke-test', metavar='FILE',
                        help='scripted GUI test run on FILE (requires a display)')
    parser.add_argument('--no-restore', action='store_true',
                        help='do not offer to restore the previous session')
    args = parser.parse_args(argv)

    for name, why in check_dependencies():
        if name in ('xlrd', 'oletools'):
            print(f"note: optional dependency '{name}' not installed — {why}")

    if args.self_test:
        if not HAS_OPENPYXL:
            print("note: openpyxl missing — .xlsx/.xlsm files will fail to load")
        return run_self_test(args.self_test)

    missing = [(n, w) for n, w in check_dependencies(for_gui=True)
               if n in ('openpyxl', 'matplotlib', 'reportlab')]
    if missing:
        for name, why in missing:
            print(f"error: missing dependency '{name}' — {why}")
        print("Install with: pip install " + " ".join(name for name, _ in missing))
        return 3

    if not TK_AVAILABLE:
        print("error: tkinter is not available — install your OS package "
              "(e.g. 'apt install python3-tk') to use the GUI, or run --self-test.")
        return 3

    matplotlib.use('TkAgg')
    global FigureCanvasTkAgg
    from matplotlib.backends.backend_tkagg import FigureCanvasTkAgg

    if args.smoke_test:
        return run_smoke_test(args.smoke_test)

    app = AccountingApp(restore_session=not args.no_restore)
    app.mainloop()
    return 0


if __name__ == '__main__':
    sys.exit(main())
