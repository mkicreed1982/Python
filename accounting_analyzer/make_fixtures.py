#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Generate test fixtures for AccountingAnalyzer: one small finance dataset
(general ledger + bank statement + P&L) written to every supported format:

  fixtures/ledger.xlsx   (openpyxl; includes a Totals sheet with an uncached formula)
  fixtures/ledger.xlsm   (same workbook with an embedded vbaProject.bin for macro detection)
  fixtures/ledger.xls    (xlwt)
  fixtures/ledger.xml    (Excel 2003 SpreadsheetML)
  fixtures/ledger.csv    (Ledger sheet only)
  fixtures/ledger.ods    (minimal OpenDocument zip)

The dataset intentionally contains DUPLICATE amounts (two rents of -500.00 in
the ledger, only one in the bank statement) to exercise the reconciliation
duplicate-matching fix, and one large outlier for anomaly detection.
"""

import os
import zipfile
from datetime import datetime
from xml.sax.saxutils import escape

FIXTURES_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'fixtures')

LEDGER_HEADERS = ['Date', 'Description', 'Category', 'Amount', 'Balance']
LEDGER_ROWS = [
    [datetime(2025, 1, 5),  'Office rent January',    'Rent',      -500.00, 9500.00],
    [datetime(2025, 1, 6),  'Warehouse rent January', 'Rent',      -500.00, 9000.00],
    [datetime(2025, 1, 8),  'Client invoice #101',    'Sales',     1250.75, 10250.75],
    [datetime(2025, 1, 10), 'Client invoice #102',    'Sales',     1250.75, 11501.50],
    [datetime(2025, 1, 12), 'Stationery order',       'Supplies',  -89.99,  11411.51],
    [datetime(2025, 1, 15), 'Electricity bill',       'Utilities', -142.50, 11269.01],
    [datetime(2025, 1, 18), 'Client invoice #103',    'Sales',     2400.00, 13669.01],
    [datetime(2025, 1, 20), 'Software subscription',  'Supplies',  -49.00,  13620.01],
    [datetime(2025, 1, 22), 'Salaries January',       'Salaries',  -3200.00, 10420.01],
    [datetime(2025, 1, 25), 'Client invoice #104',    'Sales',     980.25,  11400.26],
    [datetime(2025, 1, 27), 'Internet bill',          'Utilities', -59.90,  11340.36],
    [datetime(2025, 1, 29), 'Client contract #105',   'Sales',     50000.00, 61340.36],
    [datetime(2025, 1, 30), 'Bank fees',              'Fees',      -25.00,  61315.36],
    [datetime(2025, 1, 31), 'Client invoice #106',    'Sales',     310.40,  61625.76],
]

BANK_HEADERS = ['Date', 'Reference', 'Amount']
BANK_ROWS = [
    [datetime(2025, 1, 5),  'SEPA-99011', -500.00],   # matches ONE of the two rents
    [datetime(2025, 1, 8),  'TRN-10101',  1250.75],
    [datetime(2025, 1, 10), 'TRN-10102',  1250.75],
    [datetime(2025, 1, 18), 'TRN-10103',  2400.00],
    [datetime(2025, 1, 22), 'PAY-55010',  -3200.00],
    [datetime(2025, 1, 25), 'TRN-10104',  980.25],
    [datetime(2025, 1, 15), 'DD-77123',   -142.50],
    [datetime(2025, 1, 28), 'FEE-00021',  -12.00],    # bank-only entry
]

PNL_HEADERS = ['Month', 'Revenue', 'Expenses']
PNL_ROWS = [
    ['January',  56192.15, 4566.39],
    ['February', 43210.00, 5120.10],
    ['March',    47990.50, 4890.75],
]

SHEETS = [
    ('Ledger', LEDGER_HEADERS, LEDGER_ROWS),
    ('Bank', BANK_HEADERS, BANK_ROWS),
    ('PnL', PNL_HEADERS, PNL_ROWS),
]

# A fake vbaProject.bin (OLE magic + padding). Not a parseable VBA project —
# it exists so macro *presence* detection has something to find.
FAKE_VBA_PROJECT = b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1' + b'\x00' * 504


def write_xlsx(path):
    import openpyxl
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, headers, rows in SHEETS:
        ws = wb.create_sheet(name)
        ws.append(headers)
        for row in rows:
            ws.append(row)
    totals = wb.create_sheet('Totals')
    totals.append(['Label', 'Value'])
    # formula saved by openpyxl has no cached value → triggers the loader warning
    totals.append(['Total Amount', f'=SUM(Ledger!D2:D{len(LEDGER_ROWS) + 1})'])
    wb.save(path)


def write_xlsm(xlsx_path, path):
    with zipfile.ZipFile(xlsx_path) as src, \
            zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            dst.writestr(item, src.read(item.filename))
        dst.writestr('xl/vbaProject.bin', FAKE_VBA_PROJECT)


def write_xls(path):
    import xlwt
    wb = xlwt.Workbook()
    date_style = xlwt.easyxf(num_format_str='MM/DD/YYYY')
    for name, headers, rows in SHEETS:
        ws = wb.add_sheet(name)
        for c, h in enumerate(headers):
            ws.write(0, c, h)
        for r, row in enumerate(rows, start=1):
            for c, value in enumerate(row):
                if isinstance(value, datetime):
                    ws.write(r, c, value, date_style)
                else:
                    ws.write(r, c, value)
    wb.save(path)


def _ssml_cell(value):
    if value is None:
        return '<Cell/>'
    if isinstance(value, datetime):
        return (f'<Cell><Data ss:Type="DateTime">'
                f'{value.strftime("%Y-%m-%dT%H:%M:%S.000")}</Data></Cell>')
    if isinstance(value, bool):
        return f'<Cell><Data ss:Type="Boolean">{int(value)}</Data></Cell>'
    if isinstance(value, (int, float)):
        return f'<Cell><Data ss:Type="Number">{value}</Data></Cell>'
    return f'<Cell><Data ss:Type="String">{escape(str(value))}</Data></Cell>'


def write_spreadsheetml(path):
    parts = ['<?xml version="1.0"?>',
             '<Workbook xmlns="urn:schemas-microsoft-com:office:spreadsheet" '
             'xmlns:ss="urn:schemas-microsoft-com:office:spreadsheet">']
    for name, headers, rows in SHEETS:
        parts.append(f'<Worksheet ss:Name="{escape(name)}"><Table>')
        parts.append('<Row>' + ''.join(_ssml_cell(h) for h in headers) + '</Row>')
        for row in rows:
            parts.append('<Row>' + ''.join(_ssml_cell(v) for v in row) + '</Row>')
        parts.append('</Table></Worksheet>')
    parts.append('</Workbook>')
    with open(path, 'w', encoding='utf-8') as f:
        f.write('\n'.join(parts))


def write_csv(path):
    import csv
    with open(path, 'w', encoding='utf-8', newline='') as f:
        writer = csv.writer(f)
        writer.writerow(LEDGER_HEADERS)
        for row in LEDGER_ROWS:
            writer.writerow([v.strftime('%Y-%m-%d') if isinstance(v, datetime) else v
                             for v in row])


def _ods_cell(value):
    if value is None:
        return '<table:table-cell/>'
    if isinstance(value, datetime):
        return (f'<table:table-cell office:value-type="date" '
                f'office:date-value="{value.strftime("%Y-%m-%dT%H:%M:%S")}"/>')
    if isinstance(value, bool):
        return (f'<table:table-cell office:value-type="boolean" '
                f'office:boolean-value="{"true" if value else "false"}"/>')
    if isinstance(value, (int, float)):
        return f'<table:table-cell office:value-type="float" office:value="{value}"/>'
    return (f'<table:table-cell office:value-type="string">'
            f'<text:p>{escape(str(value))}</text:p></table:table-cell>')


def write_ods(path):
    content = ['<?xml version="1.0" encoding="UTF-8"?>',
               '<office:document-content '
               'xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
               'xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" '
               'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" '
               'office:version="1.2">',
               '<office:body><office:spreadsheet>']
    for name, headers, rows in SHEETS:
        content.append(f'<table:table table:name="{escape(name)}">')
        content.append('<table:table-row>' +
                       ''.join(_ods_cell(h) for h in headers) + '</table:table-row>')
        for row in rows:
            content.append('<table:table-row>' +
                           ''.join(_ods_cell(v) for v in row) + '</table:table-row>')
        content.append('</table:table>')
    content.append('</office:spreadsheet></office:body></office:document-content>')

    manifest = ('<?xml version="1.0" encoding="UTF-8"?>'
                '<manifest:manifest '
                'xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0" '
                'manifest:version="1.2">'
                '<manifest:file-entry manifest:full-path="/" '
                'manifest:media-type="application/vnd.oasis.opendocument.spreadsheet"/>'
                '<manifest:file-entry manifest:full-path="content.xml" '
                'manifest:media-type="text/xml"/>'
                '</manifest:manifest>')

    with zipfile.ZipFile(path, 'w') as zf:
        zf.writestr('mimetype', 'application/vnd.oasis.opendocument.spreadsheet',
                    compress_type=zipfile.ZIP_STORED)
        zf.writestr('content.xml', '\n'.join(content),
                    compress_type=zipfile.ZIP_DEFLATED)
        zf.writestr('META-INF/manifest.xml', manifest,
                    compress_type=zipfile.ZIP_DEFLATED)


def main():
    os.makedirs(FIXTURES_DIR, exist_ok=True)
    xlsx_path = os.path.join(FIXTURES_DIR, 'ledger.xlsx')
    write_xlsx(xlsx_path)
    write_xlsm(xlsx_path, os.path.join(FIXTURES_DIR, 'ledger.xlsm'))
    write_xls(os.path.join(FIXTURES_DIR, 'ledger.xls'))
    write_spreadsheetml(os.path.join(FIXTURES_DIR, 'ledger.xml'))
    write_csv(os.path.join(FIXTURES_DIR, 'ledger.csv'))
    write_ods(os.path.join(FIXTURES_DIR, 'ledger.ods'))
    for name in sorted(os.listdir(FIXTURES_DIR)):
        full = os.path.join(FIXTURES_DIR, name)
        print(f"  wrote {name} ({os.path.getsize(full)} bytes)")


if __name__ == '__main__':
    main()
